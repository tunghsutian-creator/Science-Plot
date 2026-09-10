from pathlib import Path

import pytest
from openpyxl import Workbook

from sciplot_core.data_mapping import create_data_mapping_confirmation, execute_data_mapping_proposal
from sciplot_core.data_mapping.table_choice import table_choice_snapshot, select_table, proposal_for_table_columns
from sciplot_core.data_mapping.source_mapping import _prepare_mapping_frames


def workbook_source(path: Path, *, reordered: bool = False):
    workbook = Workbook()
    workbook.active.title = "Notes"
    workbook.active.append(["Instrument export; do not edit original cells"])
    sheet = workbook.create_sheet("Measured")
    rows = [
        ["Measurement", "Original values"],
        ["Wavelength", "Absorbance", None, "Absorbance", "Wavelength", "Absorbance"],
        ["nm", "a.u.", None, "a.u.", "nm", "a.u."],
        [None, "B" if reordered else "A", None, "A" if reordered else "B", "C", "C"],
        [400, 1, None, 2, 405, 3], [450, 4, None, 6, 455, 5],
        [500, 2, None, 3, 505, 4], [550, 1, None, 1, 555, 2],
        ["End of measurement"],
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def table_selection():
    return {"sheet": "Measured", "header_rows": [1], "unit_row": 2, "sample_row": 3,
            "data_start_row": 4, "data_end_row": 8}


def test_excel_selection_keeps_original_cells_and_shared_and_paired_x(tmp_path):
    source = tmp_path / "original.xlsx"
    workbook_source(source)
    before = source.read_bytes()
    snapshot = table_choice_snapshot(source, "uvvis_spectrum")
    assert [table["sheet"] for table in snapshot["tables"]] == ["Notes", "Measured"]
    selected = select_table(snapshot, table_selection())
    assert [column["index"] for column in selected["columns"]] == list(range(6))
    assert selected["columns"][2]["header"] == ""
    request = tmp_path / "request.json"
    request.write_text(f'{{"input":"{source}","rule_id":"uvvis_spectrum","template":"curve"}}')
    proposal = proposal_for_table_columns(
        selected, pairs=[{"x_column": 0, "y_column": 1}, {"x_column": 0, "y_column": 3}, {"x_column": 4, "y_column": 5}],
        request_path=request, proposal_id="excel", created_at="2026-09-10T08:00:00+00:00")
    _, frames, _, _, _ = _prepare_mapping_frames(proposal, source_root=tmp_path)
    assert frames["series_001"].values.tolist() == [[400, 2], [450, 6], [500, 3], [550, 1]]
    assert frames["series_002"].values.tolist() == [[405, 3], [455, 5], [505, 4], [555, 2]]
    receipt = create_data_mapping_confirmation(proposal, source_root=tmp_path, request_path=request,
        output_root=tmp_path / "execution", confirmed_by="test")
    executed = execute_data_mapping_proposal(proposal, receipt, source_root=tmp_path, request_path=request, output_root=tmp_path / "execution")
    assert [item["sample_label"] for item in executed["outputs"]] == ["A", "B", "C"]
    assert source.read_bytes() == before


@pytest.mark.parametrize("change", [
    {"sheet": "Missing"}, {"data_start_row": True}, {"data_end_row": 100},
    {"data_start_row": 2}, {"header_rows": [-1]}, {"header_rows": [1, 1]},
])
def test_bad_table_selection_is_correctable(tmp_path, change):
    source = tmp_path / "original.xlsx"
    workbook_source(source)
    snapshot = table_choice_snapshot(source, "uvvis_spectrum")
    with pytest.raises(ValueError):
        select_table(snapshot, {**table_selection(), **change})
    assert select_table(snapshot, table_selection())["columns"][1]["y_eligible"]


def test_ragged_csv_description_rows_do_not_shift_original_columns(tmp_path):
    source = tmp_path / "A.csv"
    source.write_text("Instrument export\n\nWavelength (nm),Absorbance (a.u.),,Notes\n400,1,,NA\n450,2,,null\nfooter\n")
    snapshot = table_choice_snapshot(source, "uvvis_spectrum")
    selected = select_table(snapshot, {"sheet": None, "header_rows": [2], "data_start_row": 3, "data_end_row": 5})
    assert selected["rows"][0]["row_index"] == 2
    assert selected["columns"][1]["y_eligible"]


def test_multirow_header_x_sample_and_float_precision_are_preserved(tmp_path):
    from sciplot_core.data_mapping.output_files import paired_table_text
    import csv
    import io

    source = tmp_path / "instrument.csv"
    source.write_text("Instrument export\nWavelength,Absorbance\n(nm),(a.u.)\nSample-A,\n400.1234567890123,0.12345678901234567\n450.2345678901234,1.2345678901234567\n")
    before = source.read_bytes()
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), {
        "sheet": None, "header_rows": [1, 2], "sample_row": 3, "data_start_row": 4, "data_end_row": 6})
    request = tmp_path / "request.json"
    request.write_text("{}")
    proposal = proposal_for_table_columns(selected, pairs=[{"x_column": 0, "y_column": 1}],
        request_path=request, proposal_id="precision", created_at="2026-09-10T08:00:00+00:00")
    _, frames, _, _, _ = _prepare_mapping_frames(proposal, source_root=tmp_path)
    rows = list(csv.reader(io.StringIO(paired_table_text(proposal, frames))))
    assert rows[2] == ["Sample-A", "Sample-A"]
    assert rows[3:] == list(csv.reader(io.StringIO(source.read_text())))[4:]
    assert source.read_bytes() == before


def test_new_table_precision_does_not_invalidate_legacy_csv_serialization(tmp_path):
    import pandas as pd
    from sciplot_core.data_mapping.output_files import _mapped_csv_sha256, _write_mapped_csv
    from sciplot_core.foundation.file_hashing import file_sha256

    frame = pd.DataFrame({"value": [1.2345678901234567]})
    modern, legacy = tmp_path / "table.csv", tmp_path / "legacy.csv"
    _write_mapped_csv(modern, frame, full_precision=True)
    _write_mapped_csv(legacy, frame)
    assert modern.read_text() == "value\n1.2345678901234567\n"
    assert legacy.read_text() == "value\n1.23456789012346\n"
    assert file_sha256(modern) == _mapped_csv_sha256(frame, full_precision=True)
    assert file_sha256(legacy) == _mapped_csv_sha256(frame)


@pytest.mark.comprehensive
def test_native_excel_task_selection_correction_and_export(tmp_path):
    from sciplot_core.task_control import start_task, resume_task, inspect_task
    source = tmp_path / "original.xlsx"
    workbook_source(source)
    original = source.read_bytes()
    task = tmp_path / "task"
    state = start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "uvvis_spectrum", "choose_columns": True}, task_dir=task)
    assert state["status"] == "needs_input", state
    state = resume_task(task, {"expected_question_id": state["question"]["question_id"], "table_selection": {**table_selection(), "header_rows": [1, 2]}})
    with pytest.raises(ValueError, match="different samples"):
        resume_task(task, {"expected_question_id": state["question"]["question_id"], "column_mapping": {"pairs": [{"x_column": 4, "y_column": 1}]}})
    state = resume_task(task, {"expected_question_id": state["question"]["question_id"], "column_mapping": {"pairs": [{"x_column": 0, "y_column": 1}, {"x_column": 0, "y_column": 3}, {"x_column": 4, "y_column": 5}]}})
    assert state["status"] == "complete", state
    current = inspect_task(task)["current_project"]
    assert current["source"]["current"] and current["qa"]["current"] and current["delivery"]["current"]
    assert source.read_bytes() == original


@pytest.mark.comprehensive
def test_mapped_annotated_update_rechecks_reordered_columns_and_rebinds_peak(tmp_path, monkeypatch):
    import json
    from openpyxl import load_workbook
    from sciplot_core.task_control import start_task, resume_task, inspect_task
    from sciplot_core.studio_core.project_query import resolve_project_figure
    from sciplot_core.studio_core.peak_analysis import inspect_peak_candidates
    from sciplot_core.studio_core.annotation_operations import inspect_annotation_state
    from sciplot_core.foundation.file_hashing import file_sha256
    from sciplot_core import task_source_execution

    def choose(task, state, pairs):
        state = resume_task(task, {"expected_question_id": state["question"]["question_id"], "table_selection": table_selection()})
        return resume_task(task, {"expected_question_id": state["question"]["question_id"], "column_mapping": {"pairs": pairs}})

    source, replacement = tmp_path / "original.xlsx", tmp_path / "revision.xlsx"
    workbook_source(source)
    workbook_source(replacement, reordered=True)
    workbook = load_workbook(replacement)
    workbook["Measured"]["A6"] = 460
    workbook.save(replacement)
    raw_bytes = {path: path.read_bytes() for path in (source, replacement)}
    task = tmp_path / "create"
    first = start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "uvvis_spectrum", "choose_columns": True}, task_dir=task)
    first = choose(task, first, [{"x_column": 0, "y_column": 1}, {"x_column": 0, "y_column": 3}])
    assert first["status"] == "complete", first
    project = Path(first["project"])
    selected = resolve_project_figure(project, None)
    document = Path(selected["document"])
    peaks = inspect_peak_candidates(project, figure_id=selected["figure_id"], object_path="/page1/graph1/series_1",
        window={"min": 400, "max": 550, "unit": "nm"}, polarity="maximum", expected_document_sha256=file_sha256(document))
    edit_task = tmp_path / "edit"
    edit = start_task({"version": 1, "action": "edit", "project": str(project), "expected_document_sha256": file_sha256(document),
        "operations": [{"op": "add_peak_label", "id": "peakA", "candidate": peaks["candidates"][0]},
            {"op": "add_reference_line", "id": "guide", "parent_path": "/page1/graph1", "axis": "x", "value": 500, "unit": "nm"},
            {"op": "add_annotation", "id": "note", "parent_path": "/page1/graph1", "text": "Fixed note", "position": {"mode": "relative", "x": 0.05, "y": 0.95}}]}, task_dir=edit_task)
    assert edit["status"] == "needs_review", edit
    assert resume_task(edit_task, {"accept_preview": True})["status"] == "complete"
    baseline = document.read_bytes()
    update_task = tmp_path / "update"
    update = start_task({"version": 1, "action": "update_source", "project": str(project), "source": str(replacement)}, task_dir=update_task)
    assert update["status"] == "needs_input" and update["question"]["field"] == "table_selection", update
    update = choose(update_task, update, [{"x_column": 0, "y_column": 3}, {"x_column": 0, "y_column": 1}])
    assert update["status"] == "needs_input" and update["question"]["field"] == "annotation_rebinding", update
    record = next(item for item in update["question"]["evidence"] if item["id"] == "peakA")
    assert record["status"] == "moved" and record["candidates"][0]["x"] == 460
    assert record["candidates"][0]["y"] == 6 and record["candidates"][0]["sample"] == "A"
    assert document.read_bytes() == baseline
    update = resume_task(update_task, {"expected_revision_id": update["revision_id"], "annotation_choices": [{
        "figure_id": record["figure_id"], "id": "peakA", "action": "rebind", "candidate_id": record["candidates"][0]["candidate_id"], "text": "460 nm"}]})
    assert update["status"] == "needs_review", update
    original_export = task_source_execution.run_export
    monkeypatch.setattr(task_source_execution, "run_export", lambda *args: (_ for _ in ()).throw(RuntimeError("simulated export interruption")))
    update = resume_task(update_task, {"expected_revision_id": update["revision_id"], "accept_source_update": True})
    assert update["status"] == "blocked" and update["phase"] == "exporting", update
    adopted = document.read_bytes()
    monkeypatch.setattr(task_source_execution, "run_export", original_export)
    update = resume_task(update_task, {"retry": True})
    assert update["status"] == "complete", update
    assert document.read_bytes() == adopted
    annotations = {item["id"]: item for item in inspect_annotation_state(project)["annotations"]}
    assert annotations["peakA"]["peak_anchor"]["x"] == 460 and annotations["peakA"]["peak_anchor"]["sample"] == "A"
    assert annotations["guide"]["value"] == 500 and annotations["note"]["position"] == {"mode": "relative", "x": 0.05, "y": 0.95}
    spec = json.loads(Path(selected["spec"]).read_text())
    assert {item["label"]: item["y_values"] for item in spec["series"]} == {"A": [2, 6, 3, 1], "B": [1, 4, 2, 1]}
    assert all(path.read_bytes() == values for path, values in raw_bytes.items())
    current = inspect_task(update_task)["current_project"]
    assert current["source"]["current"] and current["qa"]["current"] and current["delivery"]["current"]
