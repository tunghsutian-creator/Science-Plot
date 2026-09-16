"""Independent curve ranges preserve points through mapping and native revision."""

from copy import deepcopy
import csv
from dataclasses import replace
import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from sciplot_core.data_mapping import create_data_mapping_confirmation, execute_data_mapping_proposal
from sciplot_core.data_mapping.output_files import paired_table_text
from sciplot_core.data_mapping.source_mapping import _prepare_mapping_frames
from sciplot_core.data_mapping.table_choice import proposal_for_table_columns, select_table, table_choice_snapshot
from sciplot_core.semantic_sources.registered_paired_curve_transform import resolve_registered_paired_curve_transform
from sciplot_core.materials_rules import get_rule
from sciplot_core.task_control import start_task, resume_task, inspect_task


def selection(sheet="Measured", end=7):
    return {"sheet": sheet, "header_rows": [0], "unit_row": 1, "sample_row": 2,
            "data_start_row": 3, "data_end_row": end}


def source_file(path, *, separate=False, revised=False):
    a = [[400.1234567890123, 1.12345678901234], [460 if revised else 450, 4], [500, 2], [550, 1]]
    b = [[405, 2], [455, 6], [505, 3]]
    if revised:
        a.append([600, 0.5])
        b.pop()
    book = Workbook()
    first = book.active
    first.title = "Measured"
    second = book.create_sheet("Other") if separate else first
    for label, sheet, column, values in (("A", first, 1, a), ("B", second, 1 if separate else 3, b)):
        for row, values_row in enumerate([["Wavelength", "Absorbance"], ["nm", "a.u."], [label, label], *values], 1):
            for offset, value in enumerate(values_row):
                sheet.cell(row, column + offset, value)
    book.save(path)
    return {"A": a, "B": b}


def pairs(*, separate=False, revised=False):
    return [{"x_column": 0, "y_column": 1,
             "table_selection": selection(end=8 if revised else 7)},
            {"x_column": 0 if separate else 2, "y_column": 1 if separate else 3,
             "table_selection": selection("Other" if separate else "Measured", end=5 if revised else 6)}]


def proposal(tmp_path, *, separate=False):
    source = tmp_path / "source.xlsx"
    values = source_file(source, separate=separate)
    snap = select_table(table_choice_snapshot(source, "uvvis_spectrum"), selection())
    request = tmp_path / "request.json"
    request.write_text(json.dumps({"input": str(source), "rule_id": "uvvis_spectrum", "template": "curve"}))
    selected = proposal_for_table_columns(snap, pairs=pairs(separate=separate), request_path=request,
                                         proposal_id="ranges", created_at="2026-09-10T08:00:00Z")
    return source, values, snap, request, selected


@pytest.mark.parametrize("separate", [False, True])
def test_independent_ranges_survive_composite_and_scientific_adapter(tmp_path, separate):
    source, values, _, request, selected = proposal(tmp_path, separate=separate)
    before = source.read_bytes()
    _, frames, _, _, _ = _prepare_mapping_frames(selected, source_root=tmp_path)
    assert [frame.values.tolist() for frame in frames.values()] == list(values.values())
    rows = list(csv.reader(io.StringIO(paired_table_text(selected, frames))))
    assert list(map(float, rows[-1][:2])) == [550, 1]
    assert rows[-1][2:] == ["", ""]
    root = tmp_path / "mapping"
    confirmation = create_data_mapping_confirmation(selected, source_root=tmp_path,
        request_path=request, output_root=root, confirmed_by="automated_test")
    execution = execute_data_mapping_proposal(selected, confirmation, source_root=tmp_path,
        request_path=request, output_root=root)
    assert [item["rows"] for item in execution["outputs"]] == [4, 3]
    transformed = resolve_registered_paired_curve_transform(Path(execution["effective_input"]), rule=get_rule("uvvis_spectrum"))
    assert {item.sample: [list(p) for p in item.points] for item in transformed.series} == values
    assert source.read_bytes() == before
    # A signed confirmation cannot conceal changing one curve's range.
    binding = deepcopy(selected.table_confirmation)
    binding["pairs"][1]["table_selection"]["data_end_row"] = 5
    with pytest.raises(ValueError, match="reproduce"):
        _prepare_mapping_frames(replace(selected, table_confirmation=binding), source_root=tmp_path)


@pytest.mark.parametrize("change", [
    {"table_selection": selection(end=8)}, {"table_selection": selection(end=True)},
    {"table_selection": selection("Missing")}, {"metadata_confirmations": None},
    {"x_column": True}, {"y_column": 90}, {"sort": True},
])
def test_bad_pair_region_fails_without_rewriting_source(tmp_path, change):
    source, _, snap, request, _ = proposal(tmp_path)
    before = source.read_bytes()
    selection_pairs = pairs()
    selection_pairs[1].update(change)
    with pytest.raises(ValueError):
        proposal_for_table_columns(snap, pairs=selection_pairs, request_path=request,
                                   proposal_id="bad", created_at="2026-09-10T08:00:00Z")
    assert source.read_bytes() == before


def choose(task, state, selected_pairs):
    state = resume_task(task, {"expected_question_id": state["question"]["question_id"],
                               "table_selection": selection()})
    return resume_task(task, {"expected_question_id": state["question"]["question_id"],
                              "column_mapping": {"pairs": selected_pairs}})


@pytest.mark.parametrize("unit", ["deg", "degrees", "°", "degree"])
def test_angle_spellings_preserve_values_without_conversion(tmp_path, unit):
    source = tmp_path / "xrd.csv"
    source.write_text(f"2theta ({unit}),Intensity (a.u.)\n20,1.5\n25,7\n30,2\n")
    snap = select_table(table_choice_snapshot(source, "xrd_pattern"), {
        "sheet": None, "header_rows": [0], "data_start_row": 1, "data_end_row": 4})
    assert snap["columns"][0]["x_eligible"]
    request = tmp_path / "request.json"
    request.write_text("{}")
    selected = proposal_for_table_columns(snap, pairs=[{"x_column": 0, "y_column": 1}],
        request_path=request, proposal_id="angle", created_at="2026-09-10T08:00:00Z")
    _, frames, _, _, _ = _prepare_mapping_frames(selected, source_root=tmp_path)
    assert next(iter(frames.values())).values.tolist() == [["20", "1.5"], ["25", "7"], ["30", "2"]]
    assert not selected.transformations


def test_mapping_plan_rejects_adapter_loss_even_when_labels_match():
    from sciplot_core.data_mapping.plan_binding import verify_mapping_sample_identity

    with pytest.raises(ValueError, match="point counts"):
        verify_mapping_sample_identity({"expected_sample_labels": ["A", "B"], "expected_series_points": {"A": 4, "B": 3}},
            {"output": {"series_order": ["A", "B"], "series": [{"sample": "A", "point_count": 4}, {"sample": "B", "point_count": 2}]}})


def test_sample_number_is_not_a_greek_axis_symbol(tmp_path):
    from sciplot_core.foundation.file_hashing import file_sha256
    from sciplot_core.semantic_sources.paired_curve_table_metadata import axis_match

    assert axis_match("2 θ (degree)", ("2θ",))
    assert not axis_match("FeⅠ-2", ("2θ",))
    source = tmp_path / "numbered-sample.csv"
    source.write_text("2θ (degree),FeⅠ-2\n20,1\n25,4\n30,2\n")
    declarations = [{"source_sha256": file_sha256(source), "sheet": None, "column_index": 1,
        "field": field, "value": value, "evidence": evidence} for field, value, evidence in [
            ("sample", "FeⅠ-2", {"kind": "source_cell", "sheet": None, "row_index": 0,
                "column_index": 1, "text": "FeⅠ-2"}),
            ("quantity", "Intensity", {"kind": "user_statement", "asserted_by": "automated fixture",
                "statement": "The numbered sample column records Intensity in a.u."}),
            ("unit", "a.u.", {"kind": "user_statement", "asserted_by": "automated fixture",
                "statement": "The numbered sample column records Intensity in a.u."})]]
    snap = select_table(table_choice_snapshot(source, "xrd_pattern"), {
        "sheet": None, "header_rows": [0], "data_start_row": 1, "data_end_row": 4}, declarations)
    assert snap["columns"][0]["x_eligible"]
    assert snap["columns"][1]["y_eligible"]
    assert snap["columns"][1]["sample"] == "FeⅠ-2"
    assert snap["columns"][1]["raw_metadata"]["header"] == "FeⅠ-2"


def test_disjoint_vertical_pairs_keep_distinct_original_sample_rows(tmp_path):
    source = tmp_path / "vertical.csv"
    source.write_text("Wavelength,Absorbance\nnm,a.u.\nA,A\n400,1\n450,3\n500,2\n"
                      "Wavelength,Absorbance\nnm,a.u.\nB,B\n405,2\n455,6\n")
    first = {**selection(end=6), "sheet": None}
    second = {"sheet": None, "header_rows": [6], "unit_row": 7, "sample_row": 8,
              "data_start_row": 9, "data_end_row": 11}
    snap = select_table(table_choice_snapshot(source, "uvvis_spectrum"), first)
    request = tmp_path / "request.json"
    request.write_text("{}")
    result = proposal_for_table_columns(snap, pairs=[{"x_column": 0, "y_column": 1},
        {"x_column": 0, "y_column": 1, "table_selection": second}], request_path=request,
        proposal_id="vertical", created_at="2026-09-10T08:00:00Z")
    _, frames, _, _, _ = _prepare_mapping_frames(result, source_root=tmp_path)
    assert result.sample_labels == {"series_000": "A", "series_001": "B"}
    assert [len(frame) for frame in frames.values()] == [3, 2]


def test_older_degree_confirmation_keeps_its_original_mapped_header(tmp_path):
    from sciplot_core.foundation.file_hashing import file_sha256

    source = tmp_path / "legacy.csv"
    source.write_text("2theta (degrees),Intensity (a.u.)\n20,1\n25,4\n30,2\n")
    declaration = {"source_sha256": file_sha256(source), "sheet": None, "column_index": 0,
        "field": "unit", "value": "degree", "evidence": {"kind": "source_cell", "sheet": None,
            "row_index": 0, "column_index": 0, "text": "2theta (degrees)"}}
    snap = select_table(table_choice_snapshot(source, "xrd_pattern"), {
        "sheet": None, "header_rows": [0], "data_start_row": 1, "data_end_row": 4}, [declaration])
    assert snap["columns"][0]["output_header"] == "2theta (degrees) (degree)"


def test_pair_declarations_do_not_cross_worksheets_and_conflicts_remain_blocked(tmp_path):
    from sciplot_core.foundation.file_hashing import file_sha256

    source, _, snap, request, _ = proposal(tmp_path, separate=True)
    bad = {"source_sha256": file_sha256(source), "sheet": "Measured", "column_index": 1,
           "field": "unit", "value": "nm", "evidence": {"kind": "user_statement",
           "asserted_by": "automated fault injection", "statement": "Intentional conflicting unit for validation."}}
    selected_pairs = pairs(separate=True)
    selected_pairs[1]["metadata_confirmations"] = [bad]
    with pytest.raises(ValueError, match="different worksheet"):
        proposal_for_table_columns(snap, pairs=selected_pairs, request_path=request,
            proposal_id="wrong_sheet", created_at="2026-09-10T08:00:00Z")
    bad["sheet"] = "Other"
    with pytest.raises(ValueError, match="raw_metadata_conflict"):
        proposal_for_table_columns(snap, pairs=selected_pairs, request_path=request,
            proposal_id="wrong_unit", created_at="2026-09-10T08:00:00Z")


@pytest.mark.comprehensive
def test_native_unequal_revision_cross_sheet_annotation_and_export(tmp_path):
    from sciplot_core.foundation.file_hashing import file_sha256
    from sciplot_core.studio_core.project_query import resolve_project_figure
    from sciplot_core.studio_core.peak_analysis import inspect_peak_candidates
    from sciplot_core.studio_core.annotation_operations import inspect_annotation_state

    source, revised = tmp_path / "source.xlsx", tmp_path / "revised.xlsx"
    original_values = source_file(source)
    revised_values = source_file(revised, separate=True, revised=True)
    originals = {path: path.read_bytes() for path in (source, revised)}
    task = tmp_path / "create"
    state = start_task({"version": 1, "action": "create", "source": str(source),
                       "rule_id": "uvvis_spectrum", "choose_columns": True}, task_dir=task)
    state = choose(task, state, pairs())
    assert state["status"] == "complete", state
    project = Path(state["project"])
    figure = resolve_project_figure(project, None)
    spec = json.loads(Path(figure["spec"]).read_text())
    assert {item["label"]: list(map(list, zip(item["x_values"], item["y_values"], strict=True))) for item in spec["series"]} == original_values
    document = Path(figure["document"])
    peaks = inspect_peak_candidates(project, figure_id=figure["figure_id"], object_path="/page1/graph1/series_1",
        window={"min": 400, "max": 550, "unit": "nm"}, polarity="maximum", expected_document_sha256=file_sha256(document))
    edit_task = tmp_path / "edit"
    edit = start_task({"version": 1, "action": "edit", "project": str(project),
        "expected_document_sha256": file_sha256(document), "export": False,
        "operations": [{"op": "add_peak_label", "id": "peakA", "candidate": peaks["candidates"][0]},
            {"op": "add_reference_line", "id": "guide", "parent_path": "/page1/graph1", "axis": "x", "value": 500, "unit": "nm"}]}, task_dir=edit_task)
    assert edit["status"] == "needs_review", edit
    assert resume_task(edit_task, {"accept_preview": True})["status"] == "complete"
    baseline = document.read_bytes()
    task = tmp_path / "update"
    state = start_task({"version": 1, "action": "update_source", "project": str(project), "source": str(revised)}, task_dir=task)
    state = choose(task, state, pairs(separate=True, revised=True))
    assert state["status"] == "needs_input" and state["question"]["field"] == "annotation_rebinding", state
    record = next(item for item in state["question"]["evidence"] if item["id"] == "peakA")
    assert record["candidates"][0]["x"] == 460 and record["candidates"][0]["sample"] == "A"
    assert document.read_bytes() == baseline
    state = resume_task(task, {"expected_revision_id": state["revision_id"], "annotation_choices": [{
        "figure_id": record["figure_id"], "id": "peakA", "action": "rebind",
        "candidate_id": record["candidates"][0]["candidate_id"], "text": "460 nm"}]})
    assert state["status"] == "needs_review", state
    state = resume_task(task, {"expected_revision_id": state["revision_id"], "accept_source_update": True})
    assert state["status"] == "complete", state
    spec = json.loads(Path(figure["spec"]).read_text())
    assert {item["label"]: list(map(list, zip(item["x_values"], item["y_values"], strict=True))) for item in spec["series"]} == revised_values
    annotations = {item["id"]: item for item in inspect_annotation_state(project)["annotations"]}
    assert annotations["peakA"]["peak_anchor"]["x"] == 460
    assert annotations["guide"]["value"] == 500
    current = inspect_task(task)["current_project"]
    assert all(current[key]["current"] for key in ("source", "qa", "delivery"))
    assert all(path.read_bytes() == data for path, data in originals.items())
    from test_project_source_update_native import _native

    actual = _native(document)["series"]
    assert {label: list(map(list, zip(item["x"], item["y"], strict=True))) for label, item in actual.items()} == revised_values
