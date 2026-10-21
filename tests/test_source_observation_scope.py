"""Generated physical facts distinguish table scope from scientific selection."""
from copy import deepcopy
import struct

from openpyxl import Workbook, load_workbook
import pytest

from sciplot_core.data_mapping import raw_tables
from sciplot_core.data_mapping.original_tables import source_table_catalog, source_table_snapshot
from sciplot_core.data_mapping.table_choice import select_table
from sciplot_core.foundation.file_hashing import file_sha256
from sciplot_core.source_inspection.input_errors import SourceInputError
from sciplot_core.task_errors import TaskControlError
from sciplot_core.task_semantic_plan import prepare_plan


BODY = [["Time", "Mass", "Aux"], ["s", "%", "Pa"], ["Chosen"] * 3,
        ["0.10000000000000002", "80.00000000000001", 1], ["1.0000000000000002", "79.00000000000001", 2],
        [None, None, 3], [None, None, 4]]
DISCONNECTED = [["Time", "Mass"], ["s", "%"], ["Other", "Other"], [0, 9], [], [1, 8]]


def workbook(path, tables):
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def only_cells(monkeypatch, allowed):
    """Spy the actual mapping Excel loader, also trapping excluded attachments."""
    calls = []
    original = raw_tables._read_original_cells

    def read(reference, path, *, preserve_cells=False):
        key = (path.resolve(), reference.sheet)
        assert key in allowed, f"Undeclared original-cell observation: {key}"
        calls.append(key)
        return original(reference, path, preserve_cells=preserve_cells)

    monkeypatch.setattr(raw_tables, "_read_original_cells", read)
    return calls


def reference(*, source="source:0", table="table:1", block="block:0"):
    return {"source_id": source, "table_id": table, "block_id": block}


def intent(refs=None):
    value = {"x": {"quantity": "time", "unit": "s"}, "y": {"quantity": "mass", "unit": "%"}}
    if refs is not None:
        value["selection"] = {"blocks": refs}
    return value


def test_qualified_scope_precedes_cells_preserves_original_ordinals_and_all_declared_rows(tmp_path, monkeypatch):
    from sciplot_core.plot_engine.managed_sources import load_sources
    from sciplot_core.plot_ir import create_managed_document
    from sciplot_core.plot_transforms import resolve_transforms
    from sciplot_core.task_semantic_figure import build_request
    from sciplot_core.task_semantic_selection import selection_summary

    excluded = tmp_path / "excluded.txt"
    excluded.write_text("Unsupported observation must never occur\n")
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY), ("Notes", [["note"]])])
    before = source.read_bytes(), excluded.read_bytes()
    calls = only_cells(monkeypatch, {(source, "Measured")})
    value = intent([reference(source="source:1")])
    plan = prepare_plan([excluded, source], value)
    assert plan["decisions"] == []
    selected, = plan["sources"]
    assert selected["source_id"] == "source:1"
    candidate, = selected["candidates"]
    assert candidate["candidate_id"] == "table:1" and candidate["sheet"] == "Measured"
    assert candidate["table_selection"]["data_end_row"] == len(BODY)
    block, = candidate["blocks"]
    assert block["column_indices"] == [0, 1, 2] and block["point_count"] == 4
    observation = selection_summary(plan["selection"])["observation_scope"]
    assert observation[0]["observed_sheets"] == [] and observation[0]["sha256"] == file_sha256(excluded)
    assert [(table["table_id"], table["sheet"], table["observed"]) for table in observation[1]["tables"]] == [
        ("table:0", "Summary", False), ("table:1", "Measured", True), ("table:2", "Notes", False)]
    state = {"request": {"intent": value}, "semantic_plan": plan, "semantic_choices": {},
             "intent_revision": 1, "task_dir": str(tmp_path / "task")}
    wire = build_request(state)
    document = create_managed_document(wire["template_definition"], wire["data_binding"], plot_id="scope")
    resolved = resolve_transforms(document, load_sources(document, "ordinary_xy"))
    layer = wire["template_definition"]["figure_spec"]["views"][0]["layers"][0]
    dataset = resolved["datasets"][layer["dataset_id"]]
    for name, column in (("x", 0), ("y", 1)):
        expected = [None if row[column] is None else float(row[column]) for row in BODY[3:]]
        actual = dataset["columns"][name]["values"]
        assert [None if item is None else struct.pack("!d", item) for item in actual] == [
            None if item is None else struct.pack("!d", item) for item in expected]
    assert dataset["provenance"]["sources"][0]["rows"] == [3, 4, 5, 6]
    assert dataset["provenance"]["sources"][0]["sha256"] == file_sha256(source)
    assert calls and set(calls) == {(source, "Measured")}
    assert (source.read_bytes(), excluded.read_bytes()) == before


@pytest.mark.parametrize("selection", [None, {"test_names": ["Other"]}])
def test_omitted_or_test_name_only_scope_retains_conservative_discovery(tmp_path, selection):
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY)])
    value = intent()
    if selection is not None:
        value["selection"] = selection
    with pytest.raises(SourceInputError) as error:
        prepare_plan([source], value)
    assert error.value.reason_code == "semantic_table_region_ambiguous"
    assert error.value.repair["sheet"] == "Summary" and error.value.repair["row_index"] == 4


@pytest.mark.parametrize("bad", [reference(source="source:9"), reference(table="table:9")])
def test_unknown_qualified_source_or_table_is_rejected_before_any_cells(tmp_path, monkeypatch, bad):
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY)])
    calls = only_cells(monkeypatch, set())
    with pytest.raises(TaskControlError) as error:
        prepare_plan([source], intent([reference(), bad]))
    assert error.value.reason_code == "semantic_source_selection_unmatched" and not calls


def test_selected_disconnected_table_still_fails_and_unknown_block_has_no_fallback(tmp_path, monkeypatch):
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY)])
    only_cells(monkeypatch, {(source, "Summary"), (source, "Measured")})
    with pytest.raises(SourceInputError) as error:
        prepare_plan([source], intent([reference(table="table:0")]))
    assert error.value.reason_code == "semantic_table_region_ambiguous"
    with pytest.raises(TaskControlError) as error:
        prepare_plan([source], intent([reference(block="block:9")]))
    assert error.value.reason_code == "semantic_source_selection_unmatched"


def test_multiple_qualified_tables_are_observed_without_renumbering_or_default_truncation(tmp_path, monkeypatch):
    source = workbook(tmp_path / "original.xlsx", [("First", BODY), ("Summary", DISCONNECTED), ("Last", BODY)])
    calls = only_cells(monkeypatch, {(source, "First"), (source, "Last")})
    plan = prepare_plan([source], intent([reference(table="table:2"), reference(table="table:0")]))
    assert [item["candidate_id"] for item in plan["sources"][0]["candidates"]] == ["table:0", "table:2"]
    assert plan["decisions"][0]["field"] == "tables"
    assert set(calls) == {(source, "First"), (source, "Last")}


def test_test_and_block_intersection_stays_empty_in_selected_instrument_table(tmp_path, monkeypatch):
    from test_instrument_source_regions import instrument

    rows, facts = instrument(tmp_path / "generated.csv", (3, 4))
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", rows)])
    only_cells(monkeypatch, {(source, "Measured")})
    value = intent([reference()])
    value["selection"]["test_names"] = [facts[1]["test_name"]]
    with pytest.raises(TaskControlError) as error:
        prepare_plan([source], value)
    assert error.value.reason_code == "semantic_source_selection_empty"
    assert len(error.value.repair["available_blocks"]) == 2


def test_snapshot_catalog_does_not_observe_cells_and_refresh_keeps_declared_scope(tmp_path, monkeypatch):
    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY)])
    calls = only_cells(monkeypatch, {(source, "Measured")})
    assert source_table_catalog(source)["tables"] == [
        {"table_id": "table:0", "sheet": "Summary"}, {"table_id": "table:1", "sheet": "Measured"}]
    assert calls == []
    snapshot = source_table_snapshot(source, sheets=["Measured"])
    assert snapshot["observed_sheets"] == ["Measured"] and snapshot["tables"][0]["table_id"] == "table:1"
    assert len(snapshot["catalog"]) == 2
    selection = {"sheet": "Measured", "header_rows": [0], "unit_row": 1, "sample_row": 2,
                 "data_start_row": 3, "data_end_row": len(BODY)}
    assert select_table(snapshot, selection)["table_selection"] == selection
    assert set(calls) == {(source, "Measured")}
    with pytest.raises(ValueError):
        select_table(snapshot, {**selection, "sheet": "Summary"})
    for invalid in ([], ["Missing"]):
        with pytest.raises(ValueError):
            source_table_snapshot(source, sheets=invalid)


def test_unselected_sheet_byte_change_still_invalidates_scoped_snapshot_and_cache(tmp_path, monkeypatch):
    from sciplot_core.source_tables.read_session import table_read_session

    source = workbook(tmp_path / "original.xlsx", [("Summary", DISCONNECTED), ("Measured", BODY)])
    only_cells(monkeypatch, {(source, "Measured")})
    with table_read_session():
        snapshot = source_table_snapshot(source, sheets=["Measured"])
        before = deepcopy(snapshot)
        book = load_workbook(source)
        book["Summary"]["A1"] = "Changed unselected metadata"
        book.save(source)
        with pytest.raises(ValueError, match="original evidence changed"):
            select_table(snapshot, {"sheet": "Measured", "header_rows": [0], "data_start_row": 3, "data_end_row": 5})
        assert snapshot == before and file_sha256(source) != snapshot["file_sha256"]


def test_historical_full_snapshot_shape_and_replay_are_unchanged(tmp_path):
    source = workbook(tmp_path / "original.xlsx", [("Notes", [["original"]]), ("Values", [["Time", "Mass"], [0, 9], [1, 8]])])
    expected = {"kind": "sciplot_table_choice", "version": 1, "source": str(source),
        "file_sha256": file_sha256(source), "rule_id": None,
        "tables": [{"sheet": "Notes", "row_count": 1, "column_count": 1,
                    "rows": [{"row_index": 0, "cells": ["original"]}], "preview_truncated": False},
                   {"sheet": "Values", "row_count": 3, "column_count": 2,
                    "rows": [{"row_index": 0, "cells": ["Time", "Mass"]}, {"row_index": 1, "cells": ["0", "9"]},
                             {"row_index": 2, "cells": ["1", "8"]}], "preview_truncated": False}],
        "indices": "zero_based; data_end_row is exclusive"}
    assert source_table_snapshot(source) == expected
    assert select_table(expected, {"sheet": "Values", "header_rows": [0], "data_start_row": 1, "data_end_row": 3})["table_snapshot"] == expected


def test_explicit_cross_sheet_metadata_evidence_does_not_expand_measurement_scope(tmp_path, monkeypatch):
    source = workbook(tmp_path / "original.xlsx", [("Evidence", [["Chosen sample"]]), ("Measured", [["Time", "Mass"], ["s", "%"], [0, 8], [1, 7]]),
                                                    ("Summary", DISCONNECTED)])
    calls = only_cells(monkeypatch, {(source, "Measured"), (source, "Evidence")})
    snapshot = source_table_snapshot(source, sheets=["Measured"])
    selection = {"sheet": "Measured", "header_rows": [0], "unit_row": 1, "data_start_row": 2, "data_end_row": 4}
    declaration = {"source_sha256": file_sha256(source), "sheet": "Measured", "column_index": 1,
        "field": "sample", "value": "Chosen sample", "evidence": {"kind": "source_cell", "sheet": "Evidence", "row_index": 0,
                                                                    "column_index": 0, "text": "Chosen sample"}}
    assert select_table(snapshot, selection, [declaration])["columns"][1]["sample"] == "Chosen sample"
    assert snapshot["observed_sheets"] == ["Measured"] and [table["sheet"] for table in snapshot["tables"]] == ["Measured"]
    assert set(calls) == {(source, "Measured"), (source, "Evidence")}
    declaration["evidence"]["text"] = "invented"
    with pytest.raises(ValueError, match="does not match"):
        select_table(snapshot, selection, [declaration])
