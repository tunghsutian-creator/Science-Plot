"""Scientific answers stay source-bound, reviewable, correctable and reproducible."""

from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import Workbook

from sciplot_core.data_mapping import load_data_mapping_execution
from sciplot_core.data_mapping.source_mapping import _prepare_mapping_frames
from sciplot_core.data_mapping.table_choice import proposal_for_table_columns, select_table, table_choice_snapshot
from sciplot_core.foundation.file_hashing import file_sha256
from sciplot_core.mapping_contract import DataMappingProposal
from sciplot_core.task_control import start_task, resume_task, inspect_task
from sciplot_core.task_contract import TaskControlError, task_response_schema
from sciplot_core.task_storage import load_task
from sciplot_core.task_table_region import inspect_table_region


def original(tmp_path):
    source = tmp_path / "original.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Measured"
    sheet.append(["Wavelength (nm)", "A", "B", None])
    for row in range(140):
        sheet.append([400 + row, 0.1234567890123456 + row, 2 + row, None])
    sheet.cell(141, 70, "beyond preview")
    book.save(source)
    return source


def selection():
    return {"sheet": "Measured", "header_rows": [0], "data_start_row": 1, "data_end_row": 141}


def assertion(source, field, value, column=1, evidence=None):
    return {"source_sha256": file_sha256(source), "sheet": "Measured", "column_index": column,
            "field": field, "value": value, "evidence": evidence or {
                "kind": "external_reference", "uri": "https://example.org/experiment",
                "locator": "Instrument record, spectrum A", "excerpt": "Absorbance (a.u.)"}}


def confirmations(source):
    return [assertion(source, "sample", "A", evidence={"kind": "source_cell", "sheet": "Measured",
                "row_index": 0, "column_index": 1, "text": "A"}),
            assertion(source, "quantity", "Absorbance"), assertion(source, "unit", "a.u.")]


def choose(task, state, field, value):
    return resume_task(task, {"expected_question_id": state["question"]["question_id"], field: value})


def start(tmp_path):
    source = original(tmp_path)
    task = tmp_path / "task"
    state = start_task({"version": 1, "action": "create", "source": str(source),
                       "rule_id": "uvvis_spectrum", "choose_columns": True}, task_dir=task)
    return source, task, state


def test_region_beyond_preview_is_read_only_and_preserves_original_indices(tmp_path):
    source, task, state = start(tmp_path)
    before = (task / "task.json").read_bytes()
    query = {"expected_question_id": state["question"]["question_id"], "sheet": "Measured",
             "row_start": 139, "row_end": 141, "column_start": 68, "column_end": 70}
    result = inspect_table_region(task, query)
    assert result["rows"][1] == {"row_index": 140, "cells": [
        {"column_index": 68, "text": ""}, {"column_index": 69, "text": "beyond preview"}]}
    assert result["source_sha256"] == file_sha256(source)
    assert not result["writes_performed"] and (task / "task.json").read_bytes() == before
    for change in ({"row_start": True}, {"row_end": 300}, {"column_end": 200}, {"sheet": "no"},
                   {"expected_question_id": "0" * 64}):
        with pytest.raises(TaskControlError):
            inspect_table_region(task, {**query, **change})


def test_missing_metadata_diagnostics_and_source_cell_facts_are_separate(tmp_path):
    source = original(tmp_path)
    snap = table_choice_snapshot(source, "uvvis_spectrum")
    missing = select_table(snap, selection())
    assert {item["code"] for item in missing["columns"][1]["y_rejection_reasons"]} == {
        "quantity_missing_or_mismatched", "missing_unit"}
    selected = select_table(snap, selection(), confirmations(source))
    column = selected["columns"][1]
    assert column["y_eligible"] and column["raw_metadata"] == {"header": "A", "unit": "", "sample": ""}
    assert column["header"] == "Absorbance" and column["sample"] == "A"
    assert column["numeric"]["point_count"] == 140
    assert selected["columns"][3]["numeric"]["invalid_count"] == 140


@pytest.mark.parametrize("change", [
    {"source_sha256": "0" * 64}, {"sheet": "Other"}, {"column_index": 99}, {"column_index": True},
    {"field": "data"}, {"evidence": {"kind": "user_statement", "statement": "yes"}},
    {"evidence": {"kind": "source_cell", "sheet": "Measured", "row_index": 0, "column_index": 1, "text": "edited"}},
])
def test_invalid_metadata_never_changes_current_question(tmp_path, change):
    source, task, state = start(tmp_path)
    state = choose(task, state, "table_selection", selection())
    before = (task / "task.json").read_bytes()
    values = confirmations(source)
    values[0] = {**values[0], **change}
    with pytest.raises(TaskControlError):
        choose(task, state, "metadata_confirmations", values)
    assert (task / "task.json").read_bytes() == before


def test_conflicts_stay_blocked_and_correction_invalidates_old_answers(tmp_path):
    source, task, state = start(tmp_path)
    state = choose(task, state, "table_selection", selection())
    values = confirmations(source) + [assertion(source, "unit", "s", column=0)]
    state = choose(task, state, "metadata_confirmations", values)
    assert state["question"]["evidence"]["columns"][0]["metadata_conflicts"][0]["code"] == "raw_metadata_conflict"
    with pytest.raises(TaskControlError):
        choose(task, state, "column_mapping", {"pairs": [{"x_column": 0, "y_column": 1}]})
    old = deepcopy(state)
    state = choose(task, state, "metadata_confirmations", confirmations(source))
    with pytest.raises(TaskControlError, match="旧问题"):
        choose(task, old, "metadata_confirmations", [])
    assert state["question"]["evidence"]["columns"][0]["x_eligible"]
    state = choose(task, state, "metadata_confirmations", [])
    assert not state["question"]["evidence"]["columns"][1]["y_eligible"]
    assert len(load_task(task)["scientific_choice_history"]) == 4


def test_conflicting_declarations_and_explicit_different_quantity_cannot_be_relabelled(tmp_path):
    source = original(tmp_path)
    values = confirmations(source) + [assertion(source, "unit", "nm")]
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), selection(), values)
    assert selected["columns"][1]["metadata_conflicts"][0]["code"] == "conflicting_confirmations"
    book = Workbook()
    book.active.title = "Measured"
    book.active.append(["Wavelength (nm)", "Extinction"])
    book.active.append([400, 1])
    book.save(source)
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), {**selection(), "data_end_row": 2},
                            [assertion(source, "quantity", "Absorbance"), assertion(source, "unit", "a.u.")])
    assert not selected["columns"][1]["y_eligible"]
    assert selected["columns"][1]["raw_metadata"]["header"] == "Extinction"


@pytest.mark.parametrize("value", [None, False, {}, "confirmed"])
def test_invalid_confirmation_list_is_not_a_withdrawal(tmp_path, value):
    _source, task, state = start(tmp_path)
    state = choose(task, state, "table_selection", selection())
    before = (task / "task.json").read_bytes()
    with pytest.raises(TaskControlError):
        choose(task, state, "metadata_confirmations", value)
    assert (task / "task.json").read_bytes() == before


def test_declared_new_quantity_and_scale_changes_remain_blocked(tmp_path):
    source = original(tmp_path)
    values = confirmations(source)
    values[1]["value"] = "Normalized Absorbance"
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), selection(), values)
    assert selected["columns"][1]["metadata_conflicts"][0]["code"] == "unsupported_quantity"
    values = confirmations(source) + [assertion(source, "unit", "um", column=0)]
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), selection(), values)
    assert not selected["columns"][0]["x_eligible"]
    assert selected["columns"][0]["raw_metadata"]["unit"] == "nm"


def test_mapping_replays_evidence_and_rejects_metadata_tampering(tmp_path):
    source = original(tmp_path)
    before = source.read_bytes()
    selected = select_table(table_choice_snapshot(source, "uvvis_spectrum"), selection(), confirmations(source))
    request = tmp_path / "request.json"
    request.write_text("{}")
    proposal = proposal_for_table_columns(selected, pairs=[{"x_column": 0, "y_column": 1}],
        request_path=request, proposal_id="metadata", created_at="2026-09-10T08:00:00+00:00")
    frozen = DataMappingProposal.from_dict(proposal.to_dict())
    _, frames, _, _, _ = _prepare_mapping_frames(frozen, source_root=tmp_path)
    assert frames["series_000"].iloc[0].tolist() == [400, 0.1234567890123456]
    with pytest.raises(ValueError, match="does not reproduce"):
        _prepare_mapping_frames(replace(frozen, sample_labels={"series_000": "Invented"}), source_root=tmp_path)
    mutated = deepcopy(frozen.table_confirmation)
    mutated["metadata_confirmations"][0]["evidence"]["text"] = "Invented"
    with pytest.raises(ValueError, match="cell text"):
        _prepare_mapping_frames(replace(frozen, table_confirmation=mutated), source_root=tmp_path)
    assert source.read_bytes() == before


def test_user_statement_correction_remains_attributed_and_survives_task_execution(tmp_path, monkeypatch):
    from sciplot_core import task_execution
    source, task, state = start(tmp_path)
    state = choose(task, state, "table_selection", selection())
    values = confirmations(source)
    values[2]["evidence"] = {"kind": "user_statement", "asserted_by": "data owner in test",
                             "statement": "My instrument exports absorbance in a.u."}
    values[2]["value"] = "nm"
    state = choose(task, state, "metadata_confirmations", values)
    assert not state["question"]["evidence"]["columns"][1]["y_eligible"]
    values[2]["value"] = "a.u."
    state = choose(task, state, "metadata_confirmations", values)
    monkeypatch.setattr(task_execution, "create_project", lambda *args, **kwargs: {
        "project_dir": str(tmp_path / "mocked_native"), "studio_run": {"ready_to_use": True, "failure_reason": None}})
    state = choose(task, state, "column_mapping", {"pairs": [{"x_column": 0, "y_column": 1}]})
    assert state["status"] == "complete", state
    execution = load_data_mapping_execution(state["data_mapping"]["data_mapping_execution"])
    assert execution["outputs"][0]["rows"] == 140
    saved = load_task(task)
    proposal = json.loads(Path(saved["mapping_choice"]["proposal_path"]).read_text())
    assert proposal["table_confirmation"]["metadata_confirmations"][2]["evidence"]["kind"] == "user_statement"
    assert inspect_task(task)["status"] == "complete"


def test_mcp_region_and_metadata_schemas_share_the_same_service(tmp_path):
    pytest.importorskip("mcp")
    from sciplot_core.mcp_server.schemas import tool_definitions
    from sciplot_core.mcp_server.services import invoke_owner
    from sciplot_core.task_choice_schema import table_region_schema
    source, task, state = start(tmp_path)
    tools = {item.name: item for item in tool_definitions()}
    tool = tools["sciplot_task_table_region"]
    assert tool.input_schema["properties"]["query"] == table_region_schema()
    query = {"expected_question_id": state["question"]["question_id"], "sheet": "Measured",
             "row_start": 0, "row_end": 1, "column_start": 0, "column_end": 2}
    assert invoke_owner(tool.name, {"task": str(task), "query": query}) == inspect_table_region(task, query)
    Draft202012Validator(task_response_schema()).validate({
        "expected_question_id": state["question"]["question_id"], "metadata_confirmations": confirmations(source)})


@pytest.mark.comprehensive
def test_native_metadata_confirmation_creation_and_cold_export(tmp_path):
    source, task, state = start(tmp_path)
    original_bytes = source.read_bytes()
    state = choose(task, state, "table_selection", selection())
    state = choose(task, state, "metadata_confirmations", confirmations(source))
    state = choose(task, state, "column_mapping", {"pairs": [{"x_column": 0, "y_column": 1}]})
    assert state["status"] == "complete", state
    project = Path(state["project"])
    spec = json.loads((project / "studio/spec.json").read_text())
    assert len(spec["series"]) == 1 and spec["series"][0]["label"] == "A"
    assert spec["series"][0]["x_values"] == list(range(400, 540))
    assert spec["series"][0]["y_values"][0] == 0.1234567890123456
    document = (project / "studio/document.vsz").read_bytes()
    exported = start_task({"version": 1, "action": "export", "project": str(project)}, task_dir=tmp_path / "cold-export")
    assert exported["status"] == "complete", exported
    current = inspect_task(tmp_path / "cold-export")["current_project"]
    assert all(current[key]["current"] for key in ("source", "qa", "delivery"))
    assert (project / "studio/document.vsz").read_bytes() == document
    assert source.read_bytes() == original_bytes
