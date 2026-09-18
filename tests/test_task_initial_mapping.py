"""Caller-reviewed mappings batch existing choices without bypassing evidence guards."""

from copy import deepcopy
from pathlib import Path

from jsonschema import Draft202012Validator
import pytest

from sciplot_core import task_control as control, task_execution as execution
from sciplot_core.data_mapping import load_data_mapping_execution
from sciplot_core.foundation.file_hashing import file_sha256
from sciplot_core.task_contract import TaskControlError, task_request_schema, validate_task_request
from sciplot_core.task_storage import load_task


def request_for(tmp_path):
    source = tmp_path / "original.csv"
    source.write_text("Wavelength,Absorbance,Wavelength,Absorbance\nnm,a.u.,nm,a.u.\nE0,E0,E3,E3\n400,1,405,3\n450,4,455,5\n500,2,,\n")
    selection = {"sheet": None, "header_rows": [0], "unit_row": 1, "sample_row": 2,
                 "data_start_row": 3, "data_end_row": 6}
    return {"version": 1, "action": "create", "source": str(source), "rule_id": "uvvis_spectrum",
            "mapping": {"source_sha256": file_sha256(source), "table_selection": selection,
                        "column_mapping": {"pairs": [{"x_column": 0, "y_column": 1},
                            {"x_column": 2, "y_column": 3, "table_selection": {**selection, "data_end_row": 5}}]}}}


def fake_native(monkeypatch, tmp_path):
    calls = []
    monkeypatch.setattr(execution, "create_project", lambda *a, **k: calls.append(k) or {
        "project_dir": str(tmp_path / "managed"), "studio_run": {"ready_to_use": True, "failure_reason": None}})
    monkeypatch.setattr(control, "inspect_project", lambda p: {"source": {"current": True}})
    return calls


def effective(result):
    mapped = load_data_mapping_execution(result["data_mapping"]["data_mapping_execution"])
    return Path(mapped["effective_input"]).read_text()


def test_one_call_matches_interactive_mapping_and_does_not_repeat_creation(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    before = Path(request["source"]).read_bytes()
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "batch"
    batched = control.start_task(request, task_dir=task)
    assert batched["status"] == "complete" and len(calls) == 1
    assert load_task(task)["initial_mapping_attempted"] is True
    assert control.start_task(request, task_dir=task)["status"] == "complete"
    assert control.resume_task(task, {"retry": True})["status"] == "complete"
    assert len(calls) == 1
    interactive = tmp_path / "interactive"
    state = control.start_task({k: v for k, v in request.items() if k != "mapping"} | {"choose_columns": True}, task_dir=interactive)
    for field in ("table_selection", "column_mapping"):
        state = control.resume_task(interactive, {"expected_question_id": state["question"]["question_id"], field: request["mapping"][field]})
    assert state["status"] == "complete" and len(calls) == 2
    assert effective(batched) == effective(state)
    assert "500,2,," in effective(batched)  # Independent point counts are retained.
    assert Path(request["source"]).read_bytes() == before


def test_source_digest_mismatch_blocks_before_native_creation(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    request["mapping"]["source_sha256"] = "0" * 64
    calls = fake_native(monkeypatch, tmp_path)
    state = control.start_task(request, task_dir=tmp_path / "task")
    assert state["status"] == "blocked" and state["blocker"]["reason_code"] == "source_changed"
    assert not calls and "data_mapping" not in state


@pytest.mark.parametrize("invalid", ["cross_sample", "region", "missing_unit"])
def test_invalid_batch_stops_at_correctable_question(tmp_path, monkeypatch, invalid):
    request = request_for(tmp_path)
    good = deepcopy(request["mapping"])
    if invalid == "cross_sample":
        request["mapping"]["column_mapping"] = {"x_column": 0, "y_column": 3}
    elif invalid == "region":
        request["mapping"]["table_selection"]["data_end_row"] = 300
    else:
        path = Path(request["source"])
        path.write_text(path.read_text().replace("nm,a.u.,nm,a.u.", "nm,,nm,a.u."))
        request["mapping"]["source_sha256"] = file_sha256(path)
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    state = control.start_task(request, task_dir=task)
    assert state["status"] == "needs_input" and not calls
    assert state["mapping_error"]["reason_code"] == "invalid_mapping_selection"
    assert state["question"]["question_id"]
    if invalid == "region":
        state = control.resume_task(task, {"expected_question_id": state["question"]["question_id"], "table_selection": good["table_selection"]})
    if invalid == "missing_unit":
        # Choosing the complete second pair is a correction; no guessed unit.
        good["column_mapping"] = {"pairs": [good["column_mapping"]["pairs"][1]]}
    state = control.resume_task(task, {"expected_question_id": state["question"]["question_id"], "column_mapping": good["column_mapping"]})
    assert state["status"] == "complete" and len(calls) == 1


def test_metadata_batch_uses_the_same_attributed_confirmation(tmp_path, monkeypatch):
    from test_table_metadata import original, selection, confirmations

    source = original(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "uvvis_spectrum",
        "mapping": {"source_sha256": file_sha256(source), "table_selection": selection(),
                    "metadata_confirmations": confirmations(source), "column_mapping": {"x_column": 0, "y_column": 1}}},
        task_dir=tmp_path / "task")
    assert result["status"] == "complete" and len(calls) == 1
    state = load_task(tmp_path / "task")
    assert state["mapping_choice"]["question"]["evidence"]["metadata_confirmations"] == confirmations(source)


@pytest.mark.parametrize("mutation", ["array", "bool_index", "extra", "sha"])
def test_batch_schema_and_runtime_reject_malformed_mapping(tmp_path, mutation):
    request = request_for(tmp_path)
    if mutation == "array":
        request["mapping"] = [[400, 1]]
    elif mutation == "bool_index":
        request["mapping"]["column_mapping"]["pairs"][0]["x_column"] = True
    elif mutation == "extra":
        request["mapping"]["values"] = [[400, 1]]
    else:
        request["mapping"]["source_sha256"] = "old"
    assert not Draft202012Validator(task_request_schema()).is_valid(request)
    with pytest.raises(TaskControlError, match="mapping"):
        validate_task_request(request)


def test_batch_requires_explicit_rule(tmp_path):
    request = request_for(tmp_path)
    request.pop("rule_id")
    assert not Draft202012Validator(task_request_schema()).is_valid(request)
    with pytest.raises(TaskControlError, match="rule_id"):
        validate_task_request(request)


@pytest.mark.parametrize("change", [{"rule_id": "typo"}, {"template": "unsupported"}])
def test_invalid_batch_rule_or_template_keeps_existing_rule_correction(tmp_path, monkeypatch, change):
    request = request_for(tmp_path) | change
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    state = control.start_task(request, task_dir=task)
    assert state["status"] == "needs_input" and state["question"]["field"] == "rule_id" and not calls
    state = control.resume_task(task, {"rule_id": "uvvis_spectrum"})
    assert state["status"] == "complete" and len(calls) == 1


@pytest.mark.parametrize("after_confirmation", [False, True])
def test_interrupted_batch_retry_uses_original_intent_and_saved_confirmation(tmp_path, monkeypatch, after_confirmation):
    from sciplot_core import task_initial_mapping as batch

    request = request_for(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    original = batch.accept_column_response

    def interrupted(*args):
        if after_confirmation:
            original(*args)
        raise RuntimeError("interrupted mapping batch")

    monkeypatch.setattr(batch, "accept_column_response", interrupted)
    task = tmp_path / "task"
    state = control.start_task(request, task_dir=task)
    assert state["status"] == "blocked" and not calls
    confirmation = next(task.glob("column_choices/*/confirmation.json"), None)
    digest = file_sha256(confirmation) if confirmation else None
    monkeypatch.setattr(batch, "accept_column_response", original)
    state = control.resume_task(task, {"retry": True})
    assert state["status"] == "complete" and len(calls) == 1
    assert not confirmation or file_sha256(confirmation) == digest


@pytest.mark.comprehensive
def test_public_cli_batch_exports_exact_independent_pairs_and_full_inspect(tmp_path):
    import json
    import subprocess
    from test_project_source_update_native import _native

    request = request_for(tmp_path)
    original = Path(request["source"]).read_bytes()
    path = tmp_path / "request.json"
    path.write_text(json.dumps(request))
    repo = Path(__file__).resolve().parents[1]

    def cli(*args):
        output = subprocess.run([str(repo / "skill/scripts/sciplot"), *map(str, args), "--json"],
                                check=True, capture_output=True, text=True)
        return json.loads(output.stdout)

    state = cli("task", "start", "--request", path, "--task-dir", tmp_path / "task")
    assert state["status"] == "complete", state
    assert all(state["current_project"][key]["current"] for key in ("source", "qa", "delivery"))
    full = cli("task", "inspect", state["task_dir"], "--full")
    assert "figures" in full["result"] and "figures" not in state["result"]
    assert full["data_mapping"] == state["data_mapping"]
    native = _native(Path(state["current_project"]["figures"][0]["document"]))["series"]
    assert native["E0"]["x"] == [400, 450, 500] and native["E0"]["y"] == [1, 4, 2]
    assert native["E3"]["x"] == [405, 455] and native["E3"]["y"] == [3, 5]
    assert all(Path(image).is_file() for image in state["next_step"]["images"])
    assert Path(request["source"]).read_bytes() == original


def test_pending_table_recovery_batches_metadata_region_and_pairs(tmp_path, monkeypatch):
    from test_table_metadata import original, selection, confirmations

    source = original(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task({"version": 1, "action": "create", "source": str(source),
                                 "rule_id": "uvvis_spectrum"}, task_dir=task)
    assert result["status"] == "needs_input"
    assert result["next_step"]["schema_query"]["name"] == "mapping"
    response = {"expected_question_id": result["question"]["question_id"], "mapping": {
        "source_sha256": file_sha256(source), "table_selection": selection(),
        "metadata_confirmations": confirmations(source), "column_mapping": {"x_column": 0, "y_column": 1, "label": "sample A"}}}
    result = control.resume_task(task, response)
    assert result["status"] == "complete" and len(calls) == 1
    assert result["local_timing"]["calls"] == 2
    assert "sample A,sample A" in effective(result)
    assert control.resume_task(task, response)["status"] == "complete" and len(calls) == 1


def test_bad_recovery_returns_correctable_question_then_one_good_answer(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    mapping = request.pop("mapping")
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task(request | {"choose_columns": True}, task_dir=task)
    old_id = result["question"]["question_id"]
    bad = deepcopy(mapping)
    bad["column_mapping"] = {"x_column": 0, "y_column": 3}
    result = control.resume_task(task, {"expected_question_id": old_id, "mapping": bad})
    assert result["status"] == "needs_input" and not calls
    assert result["mapping_error"]["reason_code"] == "invalid_mapping_selection"
    saved = (task / "task.json").read_bytes()
    with pytest.raises(TaskControlError, match="旧问题"):
        control.resume_task(task, {"expected_question_id": old_id, "mapping": mapping})
    assert (task / "task.json").read_bytes() == saved
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "complete" and len(calls) == 1


def test_rule_correction_and_mapping_need_only_one_reply(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    mapping = request.pop("mapping")
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task(request | {"rule_id": "unknown"}, task_dir=task)
    assert result["question"]["field"] == "rule_id"
    assert result["question"]["evidence"]["file_sha256"] == mapping["source_sha256"]
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"],
        "rule_id": "uvvis_spectrum", "mapping": mapping})
    assert result["status"] == "complete" and len(calls) == 1


@pytest.mark.parametrize("after_confirmation", [False, True])
def test_interrupted_recovery_reuses_saved_intent_and_confirmation(tmp_path, monkeypatch, after_confirmation):
    from sciplot_core import task_initial_mapping as batch

    request = request_for(tmp_path)
    mapping = request.pop("mapping")
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task(request | {"choose_columns": True}, task_dir=task)
    original = batch.accept_column_response

    def interrupted(*args):
        if after_confirmation:
            original(*args)
        raise RuntimeError("interrupted one-reply recovery")

    monkeypatch.setattr(batch, "accept_column_response", interrupted)
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "blocked" and not calls
    confirmation = next(task.glob("column_choices/*/confirmation.json"), None)
    digest = file_sha256(confirmation) if confirmation else None
    monkeypatch.setattr(batch, "accept_column_response", original)
    result = control.resume_task(task, {"retry": True})
    assert result["status"] == "complete" and len(calls) == 1
    assert not confirmation or file_sha256(confirmation) == digest


def ftir_workbook_mapping(tmp_path, *, x_unit="cm^-1"):
    from openpyxl import Workbook

    source = tmp_path / "spectra.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Spectra"
    for row in [["Original FTIR measurements"], [f"Wavenumber ({x_unit})", "original 1", "original 2 before", "original 2 after"],
                [400, 90, 70, 80], [800, 60, 50, 55], [1600, 95, 75, 85]]:
        sheet.append(row)
    book.create_sheet("Info").append(["Original %T, no baseline correction"])
    book.save(source)
    digest = file_sha256(source)
    declarations = []
    for col in (1, 3):
        header = sheet.cell(2, col+1).value
        for field, value, evidence in [
            ("sample", header, {"kind": "source_cell", "sheet": "Spectra", "row_index": 1, "column_index": col, "text": header}),
            ("quantity", "%T", {"kind": "source_cell", "sheet": "Info", "row_index": 0, "column_index": 0, "text": "Original %T, no baseline correction"}),
            ("unit", "%", {"kind": "source_cell", "sheet": "Info", "row_index": 0, "column_index": 0, "text": "Original %T, no baseline correction"}),
        ]:
            declarations.append({"source_sha256": digest, "sheet": "Spectra", "column_index": col,
                                 "field": field, "value": value, "evidence": evidence})
    selection = {"sheet": "Spectra", "header_rows": [1], "data_start_row": 2, "data_end_row": 5}
    return source, {"source_sha256": digest, "table_selection": selection,
                   "metadata_confirmations": declarations, "column_mapping": {"pairs": [
                       {"x_column": 0, "y_column": 1, "label": "batch 1"},
                       {"x_column": 0, "y_column": 3, "label": "batch 2", "table_selection": {**selection, "data_end_row": 4}},
                   ]}}


def test_ftir_layout_failure_recovers_with_source_bound_choices(tmp_path, monkeypatch):
    source, mapping = ftir_workbook_mapping(tmp_path)
    before = source.read_bytes()
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=task)
    assert result["status"] == "needs_input"
    assert result["question"]["diagnostics"]["tables"][0]["columns"][0]["numeric_count"] == 3
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "complete", result
    assert len(calls) == 1 and source.read_bytes() == before
    assert effective(result).startswith("Wavenumber,%T,Wavenumber,%T\ncm^-1,%,cm^-1,%\n")
    assert "1600,95,," in effective(result)
    saved = load_task(task)
    assert saved["mapping_choice"]["question"]["evidence"]["metadata_confirmations"] == mapping["metadata_confirmations"]


@pytest.mark.comprehensive
def test_native_one_reply_recovery_retains_units_ranges_and_original_binding(tmp_path):
    from test_project_source_update_native import _native

    source, mapping = ftir_workbook_mapping(tmp_path, x_unit="cm⁻¹")
    before = source.read_bytes()
    task = tmp_path / "task"
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=task)
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "complete", result
    assert all(result["current_project"][key]["current"] for key in ("source", "qa", "delivery"))
    native = _native(Path(result["current_project"]["figures"][0]["document"]))["series"]
    assert native["batch 1"]["x"] == [400, 800, 1600]
    assert native["batch 2"]["x"] == [400, 800]
    # Stacked FTIR applies a constant display offset per curve; no rescaling.
    for label, raw in (("batch 1", [90, 60, 95]), ("batch 2", [80, 55])):
        shown = native[label]["y"]
        offset = shown[0] - raw[0]
        assert shown == pytest.approx([value + offset for value in raw], abs=1e-12)
    import csv

    delivery = Path(result["current_project"]["delivery"]["path"])
    rows = list(csv.reader(next((delivery / 'data').glob('*.csv')).open()))
    assert rows[1] == ["cm^-1", "%", "cm^-1", "%"]
    assert [[float(v) if v else None for v in row] for row in rows[3:]] == [
        [400, 90, 400, 80], [800, 60, 800, 55], [1600, 95, None, None]]
    exported = control.start_task({"version": 1, "action": "export", "project": result["project"]}, task_dir=tmp_path / "export-task")
    assert exported["status"] == "complete", exported
    assert source.read_bytes() == before
    revision_dir = tmp_path / 'revision'
    revision_dir.mkdir()
    revised_source, revised_mapping = ftir_workbook_mapping(revision_dir)
    from openpyxl import load_workbook

    workbook = load_workbook(revised_source)
    workbook['Spectra']['B3'] = 92
    workbook.save(revised_source)
    workbook.close()
    digest = file_sha256(revised_source)
    revised_mapping['source_sha256'] = digest
    for declaration in revised_mapping['metadata_confirmations']:
        declaration['source_sha256'] = digest
    update_task = tmp_path / 'update'
    update = control.start_task({'version': 1, 'action': 'update_source', 'source': str(revised_source),
                                'project': result['project']}, task_dir=update_task)
    update = control.resume_task(update_task, {'expected_question_id': update['question']['question_id'], 'mapping': revised_mapping})
    assert update['status'] == 'needs_review', update  # Mapping never accepts a scientific revision.
    assert {item['scope'] for item in update['previews']} == {'before', 'candidate'}
    assert _native(Path(result['current_project']['figures'][0]['document']))['series'] == native
    cancelled = control.resume_task(update_task, {'expected_revision_id': update['revision_id'], 'accept_source_update': False})
    assert cancelled['status'] == 'cancelled'
    source.write_bytes(before + b'changed')
    current = control.inspect_task(task)["current_project"]
    assert current["source"]["current"] is False


def test_unambiguous_metadata_layout_is_repaired_without_ai_roundtrip(tmp_path, monkeypatch):
    from openpyxl import Workbook

    source = tmp_path / "instrument.xlsx"
    book = Workbook()
    sheet = book.active
    for row in [["Instrument export"], ["Wavenumber", "Transmittance"], ["cm^-1", "%"],
                ["A", "A"], [400, 91], [500, 82], [800, 73]]:
        sheet.append(row)
    book.save(source)
    before = source.read_bytes()
    calls = fake_native(monkeypatch, tmp_path)
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=tmp_path / "task")
    assert result["status"] == "complete", result
    assert result["automatic_repairs"][0]["code"] == "explicit_metadata_layout"
    assert result["local_timing"]["calls"] == 1 and len(calls) == 1
    assert source.read_bytes() == before


@pytest.mark.parametrize("mutation", ["hole", "unit", "sample", "other_sheet"])
def test_auto_repair_never_guesses_or_drops_rows(tmp_path, monkeypatch, mutation):
    from openpyxl import Workbook

    source = tmp_path / "instrument.xlsx"
    book = Workbook()
    sheet = book.active
    for row in [["Instrument export"], ["Wavenumber", "Transmittance"], ["cm^-1", "%"],
                ["A", "A"], [400, 91], [500, 82], [800, 73]]:
        sheet.append(row)
    if mutation == "hole":
        sheet['B6'] = None
    elif mutation == "unit":
        sheet['B3'] = None
    elif mutation == "sample":
        sheet['B4'] = 'B'
    else:
        book.copy_worksheet(sheet)
    book.save(source)
    calls = fake_native(monkeypatch, tmp_path)
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=tmp_path / "task")
    assert result["status"] == "needs_input" and not calls and not result.get("automatic_repairs")


@pytest.mark.parametrize("mutation", ["stale", "invented", "conflicting_quantity", "duplicate_label"])
def test_recovery_rejects_stale_or_conflicting_scientific_choices(tmp_path, monkeypatch, mutation):
    source, mapping = ftir_workbook_mapping(tmp_path)
    if mutation == "stale":
        mapping["source_sha256"] = "0" * 64
    elif mutation == "invented":
        mapping["metadata_confirmations"][1]["evidence"]["text"] = "Transmittance %"
    elif mutation == "conflicting_quantity":
        mapping["metadata_confirmations"].append({**mapping["metadata_confirmations"][1], "value": "Absorbance",
            "evidence": {"kind": "external_reference", "uri": "https://example.org", "locator": "conflicting record", "excerpt": "Absorbance"}})
    else:
        mapping["column_mapping"]["pairs"][1]["label"] = "batch 1"
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=task)
    result = control.resume_task(task, {"expected_question_id": result["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "needs_input" and not calls


def test_response_schema_rejects_unadvertised_recovery_fields(tmp_path, monkeypatch):
    from sciplot_core.task_contract import task_response_schema

    request = request_for(tmp_path)
    mapping = request.pop("mapping")
    fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    result = control.start_task(request | {"choose_columns": True}, task_dir=task)
    response = {"expected_question_id": result["question"]["question_id"], "mapping": mapping}
    schema = Draft202012Validator(task_response_schema())
    assert schema.is_valid(response)
    bad = {**response, "values": [[1, 2]]}
    assert not schema.is_valid(bad)
    before = (task / 'task.json').read_bytes()
    with pytest.raises(TaskControlError):
        control.resume_task(task, bad)
    assert (task / 'task.json').read_bytes() == before


def test_changed_original_during_recovery_blocks_without_creation(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    mapping = request.pop('mapping')
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / 'task'
    result = control.start_task(request | {'choose_columns': True}, task_dir=task)
    source = Path(request['source'])
    source.write_text(source.read_text().replace('400,1,', '400,9,'))
    result = control.resume_task(task, {'expected_question_id': result['question']['question_id'], 'mapping': mapping})
    assert result['status'] == 'blocked' and result['blocker']['reason_code'] == 'source_changed'
    assert not calls


def test_mapping_cannot_be_used_as_an_edit_response(tmp_path):
    state = {'status': 'needs_review', 'phase': 'previewing', 'request': {'action': 'edit'}}
    with pytest.raises(TaskControlError) as error:
        control._resume(tmp_path, state, {'mapping': {}})
    assert error.value.reason_code == 'invalid_task_response'


def test_legacy_column_answer_can_name_its_display_label(tmp_path, monkeypatch):
    request = request_for(tmp_path)
    request.pop('mapping')
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / 'task'
    result = control.start_task(request | {'choose_columns': True}, task_dir=task)
    result = control.resume_task(task, {'expected_question_id': result['question']['question_id'],
        'column_mapping': {'x_column': 0, 'y_column': 1, 'label': 'renamed E0'}})
    assert result['status'] == 'complete' and len(calls) == 1
    from sciplot_core.data_mapping import load_data_mapping_execution

    mapped = load_data_mapping_execution(result['data_mapping']['data_mapping_execution'])
    assert [item['sample_label'] for item in mapped['outputs']] == ['renamed E0']


def test_metadata_recovery_reports_each_bad_declaration_with_exact_target(tmp_path, monkeypatch):
    source, mapping = ftir_workbook_mapping(tmp_path)
    fake_native(monkeypatch, tmp_path)
    task = tmp_path / "task"
    state = control.start_task({"version": 1, "action": "create", "source": str(source), "rule_id": "ftir_spectrum"}, task_dir=task)
    assert "quantity='%T', unit='%'" in state["question"]["metadata_hint"]
    mapping["metadata_confirmations"][0]["sheet"] = "Info"
    mapping["metadata_confirmations"][1]["value"] = "Spectral response"
    result = control.resume_task(task, {"expected_question_id": state["question"]["question_id"], "mapping": mapping})
    assert result["status"] == "needs_input"
    message = result["mapping_error"]["message"]
    assert "2 invalid metadata declaration(s)" in message
    assert "metadata_confirmations[0] target sheet='Info' column=1 field=sample" in message
    assert "Target worksheet must be 'Spectra'" in message
    assert "metadata_confirmations[1] target sheet='Spectra' column=1 field=quantity" in message
    assert "'Spectral response' is absent" in message
    assert "Original %T" in message
    assert result["question"]["metadata_hint"] == state["question"]["metadata_hint"]


def test_source_cell_unit_spelling_equivalence_preserves_original_evidence(tmp_path):
    from openpyxl import load_workbook
    from sciplot_core.data_mapping.table_choice import table_choice_snapshot, select_table

    source, mapping = ftir_workbook_mapping(tmp_path)
    book = load_workbook(source)
    book["Spectra"].cell(2, 1).value = "Wavenumber (cm⁻¹)"
    book.save(source)
    digest = file_sha256(source)
    declarations = [{**item, "source_sha256": digest} for item in mapping["metadata_confirmations"]]
    declaration = {"source_sha256": digest, "sheet": "Spectra", "column_index": 0, "field": "unit", "value": "cm^-1",
                   "evidence": {"kind": "source_cell", "sheet": "Spectra", "row_index": 1, "column_index": 0,
                                "text": "Wavenumber (cm⁻¹)"}}
    snapshot = table_choice_snapshot(source, "ftir_spectrum")
    result = select_table(snapshot, mapping["table_selection"], [declaration, *declarations])
    assert result["columns"][0]["x_eligible"]
    assert result["columns"][0]["raw_metadata"]["header"] == "Wavenumber (cm⁻¹)"
    assert result["metadata_confirmations"][0] == declaration
    with pytest.raises(ValueError, match="Value 'nm' is absent"):
        select_table(snapshot, mapping["table_selection"], [{**declaration, "value": "nm"}, *declarations])
    with pytest.raises(ValueError, match="8 further invalid declarations omitted") as exc:
        select_table(snapshot, mapping["table_selection"], [{**declaration, "value": "nm"}] * 16)
    assert "metadata_confirmations[7]" in str(exc.value)
    assert "metadata_confirmations[8]" not in str(exc.value)


def wide_ftir_source(tmp_path, *, note='Original %T', hole=False):
    from openpyxl import Workbook

    source = tmp_path / 'wide.xlsx'
    book = Workbook()
    book.active.title = 'Data'
    for row in [['Wavenumber (cm⁻¹)', '1', '2'], [4000, 90, 95], [2000, 50, None if hole else 60], [400, 0, 0]]:
        book.active.append(row)
    book.create_sheet('Notes').append([note])
    book.save(source)
    return source


def test_candidate_reply_avoids_metadata_transcription_and_selects_ordered_subset(tmp_path, monkeypatch):
    source = wide_ftir_source(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / 'task'
    state = control.start_task({'version': 1, 'action': 'create', 'source': str(source), 'rule_id': 'ftir_spectrum'}, task_dir=task)
    assert state['status'] == 'needs_input' and not calls
    candidate, = state['question']['mapping_candidates']
    assert 'mapping' not in candidate
    assert candidate['x']['column'] == 0
    assert candidate['y']['quantity'] == '%T' and candidate['y']['unit'] == '%'
    assert candidate['points_per_pair'] == 3
    assert state['next_step']['schema_query']['name'] == 'mapping_candidate_id'
    assert [p['sample'] for p in candidate['pairs']] == ['1', '2']
    assert load_task(task)['question']['mapping_candidates'][0]['mapping']['metadata_confirmations']
    reply = {'expected_question_id': state['question']['question_id'], 'mapping_candidate_id': candidate['candidate_id'], 'pair_indices': [1, 0]}
    assert Draft202012Validator(control.task_response_schema()).is_valid(reply)
    result = control.resume_task(task, reply)
    assert result['status'] == 'complete', result
    assert len(calls) == 1
    text = effective(result)
    assert text.startswith('Wavenumber,%T,Wavenumber,%T\ncm⁻¹,%,cm⁻¹,%\n')
    assert '4000,95,4000,90' in text
    assert '400,0,400,0' in text


@pytest.mark.parametrize('reply_change,reason', [
    ({'expected_question_id': '0'*64}, 'stale_task_question'),
    ({'mapping_candidate_id': '0'*64}, 'invalid_mapping_candidate'),
    ({'pair_indices': [2]}, 'invalid_mapping_candidate'),
    ({'pair_indices': [0, 0]}, 'invalid_task_response'),
])
def test_candidate_rejections_preserve_question_and_create_nothing(tmp_path, monkeypatch, reply_change, reason):
    source = wide_ftir_source(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path / 'task'
    state = control.start_task({'version': 1, 'action': 'create', 'source': str(source), 'rule_id': 'ftir_spectrum'}, task_dir=task)
    before = (task/'task.json').read_bytes()
    reply = {'expected_question_id': state['question']['question_id'], 'mapping_candidate_id': state['question']['mapping_candidates'][0]['candidate_id'], **reply_change}
    with pytest.raises(TaskControlError) as exc:
        control.resume_task(task, reply)
    assert exc.value.reason_code == reason
    assert (task/'task.json').read_bytes() == before
    assert not calls


@pytest.mark.parametrize('note,hole', [('No response metadata', False), ('Absorbance and %T', False), ('Original %T', True)])
def test_candidates_refuse_missing_conflicting_metadata_and_measurement_holes(tmp_path, note, hole):
    from sciplot_core.task_column_mapping import table_question

    source = wide_ftir_source(tmp_path, note=note, hole=hole)
    question = table_question(source, {'rule_id': 'ftir_spectrum'})
    assert not question.get('mapping_candidates')


def test_candidate_source_change_is_blocked_without_native_creation(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    source = wide_ftir_source(tmp_path)
    calls = fake_native(monkeypatch, tmp_path)
    task = tmp_path/'task'
    state = control.start_task({'version': 1, 'action': 'create', 'source': str(source), 'rule_id': 'ftir_spectrum'}, task_dir=task)
    book = load_workbook(source)
    book['Data']['B2'] = 123
    book.save(source)
    result = control.resume_task(task, {'expected_question_id': state['question']['question_id'], 'mapping_candidate_id': state['question']['mapping_candidates'][0]['candidate_id']})
    assert result['status'] == 'blocked'
    assert result['blocker']['reason_code'] == 'source_changed'
    assert not calls


def test_invalid_request_names_only_bad_fields_and_retains_supported_output_hint():
    with pytest.raises(TaskControlError) as exc:
        validate_task_request({'version': 1, 'action': 'create', 'source': '/original.xlsx',
                               'task_path': '/task', 'out': '/delivery'})
    assert exc.value.reason_code == 'invalid_task_fields'
    assert "不支持字段：['task_path']" in str(exc.value)
    assert "缺少字段：[]" in str(exc.value)
    assert "如 out" in str(exc.value)
