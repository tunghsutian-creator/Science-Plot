"""Public generated evidence, pre-reservation rejection and saved exact replay."""

from copy import deepcopy
import json

import pytest
from openpyxl import Workbook, load_workbook
from test_managed_description_scopes import _service
from test_managed_source_reconfirmation import basic, fresh_value, patch

from sciplot_core.foundation.file_hashing import file_sha256
from sciplot_core.plot_document import seal_document
from sciplot_core.plot_engine import current_unit_admission, managed_state
from sciplot_core.project_lifecycle import publication_repair
from sciplot_core.project_lifecycle import application as service
from sciplot_core.plot_engine.content_store import intern_document
from sciplot_core.plot_engine.current_unit_admission import admit_current_units, prepare_current_update
from sciplot_core.plot_engine.managed_sources import load_sources
from sciplot_core.plot_engine.managed_updates import prepare_update
from sciplot_core.plot_engine.service import PlotService
from sciplot_core.plot_engine.storage import digest, write


def forbidden(*args, **kwargs):
    raise AssertionError("An effect ran before source closure admission or during saved replay")


def workbook_document(tmp_path):
    document = basic(tmp_path)
    book = Workbook()
    data = book.active
    data.title = "Measurements"
    for row in [("Frequency", "Modulus"), ("Hz", "Pa"), (0, 1), (5, None), (10, -3)]:
        data.append(row)
    book.create_sheet("Metadata").append(["Confirmed sample Modulus"])
    path = tmp_path / "generated.xlsx"
    book.save(path)
    book.close()
    source = document["scientific"]["data_sources"][0]
    source.update(path=str(path), sha256=file_sha256(path))
    source["table_selection"]["sheet"] = "Measurements"
    source["metadata_confirmations"] = [{"source_sha256": source["sha256"], "sheet": "Measurements",
        "column_index": 1, "field": "sample", "value": "Modulus", "evidence": {
            "kind": "source_cell", "sheet": "Metadata", "row_index": 0, "column_index": 0, "text": "Confirmed sample Modulus"}}]
    return seal_document(document), path


def test_cross_sheet_evidence_reconfirmed_from_same_workbook_and_fresh_sha(tmp_path):
    document, path = workbook_document(tmp_path)
    assert load_sources(document, "uvvis_spectrum")["raw:A"]["columns"]["column:1"]["values"] == [1, None, -3]
    book = load_workbook(path)
    book["Measurements"].cell(3, 2).value = 8
    book.save(path)
    book.close()
    updated = prepare_current_update(tmp_path, document, patch(document, fresh_value(document, path)))[0]
    loaded = load_sources(json.loads(json.dumps(updated)), "uvvis_spectrum")
    assert loaded["raw:A"]["columns"]["column:1"]["values"] == [8, None, -3]
    assert loaded["raw:A"]["provenance"]["sources"][0]["rows"] == [2, 3, 4]
    assert loaded["raw:A"]["provenance"]["sources"][0]["table_selection"]["sheet"] == "Measurements"


@pytest.mark.parametrize("change", ["metadata-proof", "source-unit", "missing-group"])
def test_invalid_fresh_source_closure_has_no_reservation_journal_executor_or_head_change(tmp_path, monkeypatch, change):
    document, path = workbook_document(tmp_path)
    if change == "missing-group":
        document["scientific"]["data_sources"].append({**deepcopy(document["scientific"]["data_sources"][0]), "source_id": "raw:B"})
        document = seal_document(document)
    book = load_workbook(path)
    if change == "metadata-proof":
        book["Metadata"].cell(1, 1).value = "Confirmed sample Other"
    elif change == "source-unit":
        book["Measurements"].cell(2, 2).value = "kPa"
    else:
        book["Measurements"].cell(3, 2).value = 9
    book.save(path)
    book.close()
    value = fresh_value(document, path)
    if change == "missing-group":
        value["bindings"].pop("raw:B")
    head = {"document": document, "binding": {}}
    original = deepcopy(head)
    monkeypatch.setattr(service, "load_head", lambda root: head)
    monkeypatch.setattr(service, "require_current", forbidden)
    monkeypatch.setattr(managed_state, "scientific_binding", forbidden)
    monkeypatch.setattr(publication_repair, "reserve_repair", forbidden)
    monkeypatch.setattr(service, "persist", forbidden)
    monkeypatch.setattr(service, "advance", forbidden)
    root = tmp_path / "unreserved"
    with pytest.raises(ValueError):
        PlotService(backend=object())._patch(root, patch(document, value))
    assert head == original and not root.exists()


def test_stale_unsaved_revision_rejects_before_source_read_or_reservation(tmp_path, monkeypatch):
    document = basic(tmp_path)
    request = patch(document, "a" * 64)
    request["base_revision"] = 8
    monkeypatch.setattr(service, "load_head", lambda root: {"document": document, "binding": {}})
    monkeypatch.setattr(current_unit_admission, "admit_current_units", forbidden)
    monkeypatch.setattr(service, "advance", forbidden)
    with pytest.raises(ValueError) as error:
        PlotService(backend=object())._patch(tmp_path / "store", request)
    assert error.value.reason_code == "document_revision_conflict"


def test_lost_reply_and_cold_service_old_saved_key_replay_exact_request_before_current_schema(tmp_path, monkeypatch):
    document = basic(tmp_path)
    # This historical scalar is no longer safe for a fresh multi-binding request.
    request = patch(document, "a" * 64, key="already-complete")
    tx = {"phase": "complete", "request": request, "request_hash": digest(request), "result": "original exact outcome"}
    saved = tmp_path / "saved.json"
    write(saved, tx)
    (tmp_path / "figure.csv").unlink()  # A historical replay cannot require current original input bytes.
    monkeypatch.setattr(service, "transaction_path", lambda *args: saved)
    monkeypatch.setattr(service, "load_head", forbidden)
    monkeypatch.setattr(current_unit_admission, "admit_current_units", forbidden)
    monkeypatch.setattr(service, "replay", lambda root, value, backend: value)
    for _lost_reply_or_new_instance in range(2):
        assert PlotService(backend=object())._patch(tmp_path, deepcopy(request)) == tx
    changed = deepcopy(request)
    changed["changes"][0]["value"] = "b" * 64
    with pytest.raises(ValueError) as error:
        PlotService(backend=object())._patch(tmp_path, changed)
    assert error.value.reason_code == "document_idempotency_conflict"


@pytest.mark.parametrize("figure", [True, False])
def test_full_summary_edit_share_virtual_source_contract_and_keep_presentation_unchanged(tmp_path, monkeypatch, figure):
    owner, root, document = _service(tmp_path, monkeypatch, figure=figure)
    original = deepcopy(document)
    target = "source:" + document["scientific"]["data_sources"][0]["source_id"]
    full = owner.describe(root)
    summary = owner.describe(root, description={"scope": "summary"})
    edit = owner.describe(root, description={"scope": "edit", "target": target, "property": "source.sha256"})
    contract = "figure_edit_contract" if figure else "object_edit_contract"
    assert edit[contract][target]["source.sha256"] == full[contract][target]["source.sha256"]
    assert summary["edit_targets"][target] == {"kind": "source", "label": target.removeprefix("source:"), "properties": ["source.sha256"]}
    assert edit["current_value"]["sha256"] == document["scientific"]["data_sources"][0]["sha256"]
    assert document == original and target not in document["presentation"]["objects"]


def test_cas_binding_refresh_changes_only_affected_leaf_and_json_reopens(tmp_path):
    from test_managed_source_reconfirmation import context_for
    from test_current_unit_admission import expert_wire, managed
    document = managed(expert_wire(tmp_path, header="sensor_C"))
    source = document["scientific"]["data_sources"][0]
    document["scientific"]["provenance"]["binding"]["semantic_task"] = {"resolution": {"bindings": [context_for(document, source["sha256"])]}}
    document["scientific"]["provenance"]["datasets"] = {"unrelated": list(range(100))}
    document = intern_document(tmp_path, seal_document(document), minimum_bytes=2)
    before = deepcopy(document)
    path = tmp_path / "figure.csv"
    path.write_text(path.read_text().replace("1,21", "1,24"))
    updated = prepare_update(document, patch(document, fresh_value(document, path, root=tmp_path)), root=tmp_path)[0]
    assert updated["scientific"]["provenance"]["datasets"] == before["scientific"]["provenance"]["datasets"]
    assert before["scientific"]["provenance"]["binding"]["kind"] == "sciplot_content_reference"
    assert "kind" not in updated["scientific"]["provenance"]["binding"]
    admit_current_units(tmp_path, json.loads(json.dumps(updated)), previous=document)


def test_virtual_source_namespace_collision_preserves_existing_object_operations(tmp_path, monkeypatch):
    from sciplot_core.plot_ir.figure_document import figure_spec, store_spec
    owner, root, document = _service(tmp_path, monkeypatch)
    spec = figure_spec(document)
    spec["id"] = "source:raw:A"
    service.load_head(root)["document"] = store_spec(document, spec)
    summary = owner.describe(root, description={"scope": "summary"})
    assert summary["edit_targets"]["source:raw:A"]["kind"] == "figure"
    assert set(summary["edit_targets"]["source:raw:A"]["properties"]) == {"theme", "layout", "source.sha256"}
    for prop in ("theme", "source.sha256"):
        edit = owner.describe(root, description={"scope": "edit", "target": "source:raw:A", "property": prop})
        assert prop in edit["figure_edit_contract"]["source:raw:A"]
