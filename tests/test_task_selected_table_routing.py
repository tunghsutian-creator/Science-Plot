"""Only the chosen original table determines its semantic builder."""
import csv
import hashlib
import math
from io import StringIO

import pytest
from openpyxl import Workbook, load_workbook

from sciplot_core.task_errors import TaskControlError
from sciplot_core.task_semantic_creation import build_managed_request
from sciplot_core.task_semantic_plan import prepare_plan
from test_frequency_interval_binding import instrument_source, intent


def mixed_book(tmp_path):
    text = instrument_source(tmp_path / "instrument.csv", temperatures=(190,))
    book = tmp_path / "alternatives.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Instrument"
    for line in text.read_text(encoding="utf-16").splitlines():
        cells = []
        for value in line.split("\t"):
            try:
                cells.append(float(value))
            except ValueError:
                cells.append(value)
        sheet.append(cells)
    sheet = workbook.create_sheet("Canonical")
    for row in [["Angular Frequency", "Complex Viscosity"], ["rad/s", "Pa s"], [1, 25], [10, 50]]:
        sheet.append(row)
    workbook.save(book)
    return book


def state(sources, tmp_path, tables):
    return {"request": {"intent": intent()}, "semantic_plan": prepare_plan(sources, intent()),
            "semantic_choices": {"tables": tables}, "intent_revision": 1, "task_dir": str(tmp_path / "task")}


@pytest.mark.parametrize("table,sheet,count", [("table:1", "Canonical", 1), ("table:0", "Instrument", 1)])
def test_selected_table_chooses_its_actual_existing_builder(tmp_path, table, sheet, count):
    book = mixed_book(tmp_path)
    original = book.read_bytes()
    wire = build_managed_request(state([book], tmp_path, {"source:0": table}))
    assert len(wire["data_binding"]["data_sources"]) == count
    assert {binding["table_selection"]["sheet"] for binding in wire["data_binding"]["data_sources"]} == {sheet}
    assert book.read_bytes() == original


def test_actual_mixed_selected_formats_bind_independent_original_science(tmp_path):
    from sciplot_core.plot_engine.managed_sources import load_sources
    from sciplot_core.plot_ir import create_managed_document
    from sciplot_core.plot_transforms import resolve_transforms

    book = mixed_book(tmp_path)
    other = instrument_source(tmp_path / "other.csv", temperatures=(150,))
    originals = {path: path.read_bytes() for path in (book, other)}
    # Read the source cells independently of the production source scanner.
    workbook = load_workbook(book, read_only=True, data_only=True)
    try:
        canonical = list(workbook["Canonical"].values)
    finally:
        workbook.close()
    instrument = list(csv.reader(StringIO(originals[other].decode("utf-16")), delimiter="\t"))
    assert canonical[:2] == [("Angular Frequency", "Complex Viscosity"), ("rad/s", "Pa s")]
    assert instrument[7][3:6] == ["Angular Frequency", "Storage Modulus", "Loss Modulus"]
    assert instrument[7][11] == "Im(Complex Viscosity)"
    current = state([book, other], tmp_path, {"source:0": "table:1"})
    original_intent = intent()
    wire = build_managed_request(current)
    assert current["request"]["intent"] == original_intent
    assert wire["data_binding"]["provenance"]["semantic_task"]["intent"] == original_intent
    sources = {binding["path"]: binding for binding in wire["data_binding"]["data_sources"]}
    assert len(wire["data_binding"]["data_sources"]) == len(sources) == 2
    expected_regions = {str(book): {"sheet": "Canonical", "header_rows": [0], "unit_row": 1,
                                   "data_start_row": 2, "data_end_row": 4},
                        str(other): {"sheet": None, "header_rows": [7], "unit_row": 9,
                                    "data_start_row": 10, "data_end_row": 26}}
    for path, original in originals.items():
        assert sources[str(path)]["sha256"] == hashlib.sha256(original).hexdigest()
        assert sources[str(path)]["table_selection"] == expected_regions[str(path)]
    bound = current["semantic_plan"]["resolved_bindings"]
    assert len(bound) == 2 and bound[0]["source_path"] == str(book) and bound[1]["source_path"] == str(other)
    assert bound[0]["coordinates"]["y"]["column_id"] == "column:1"
    assert "derivation" not in bound[0]["coordinates"]["y"]
    assert {role: column["column_id"] for role, column in bound[1]["coordinates"]["y"]["derivation"]["inputs"].items()} == {
        "angular_frequency": "column:3", "storage_modulus": "column:4", "loss_modulus": "column:5"}
    document = create_managed_document(wire["template_definition"], wire["data_binding"], plot_id="mixed-selected-originals")
    datasets = resolve_transforms(document, load_sources(document, original_intent["rule_id"]))["datasets"]
    layers = wire["template_definition"]["figure_spec"]["views"][0]["layers"]
    assert len(layers) == 2
    canonical_dataset, instrument_dataset = [datasets[layer["dataset_id"]] for layer in layers]
    for role, index in (("x", 0), ("y", 1)):
        assert [float(value).hex() for value in canonical_dataset["columns"][role]["values"]] == [
            float(row[index]).hex() for row in canonical[2:4]]
    assert [float(value).hex() for value in instrument_dataset["columns"]["x"]["values"]] == [
        float(row[3]).hex() for row in instrument[10:26]]
    assert instrument_dataset["columns"]["y"]["values"] == pytest.approx([
        math.hypot(float(row[4]), float(row[5])) / float(row[3]) for row in instrument[10:26]])
    assert instrument_dataset["columns"]["y"]["values"] != [float(row[11]) * .001 for row in instrument[10:26]]
    for dataset, path, rows, derived in ((canonical_dataset, book, [2, 3], False),
                                         (instrument_dataset, other, list(range(10, 26)), True)):
        origin = dataset["provenance"]["sources"][0]
        assert origin["sha256"] == hashlib.sha256(originals[path]).hexdigest()
        assert origin["table_selection"] == expected_regions[str(path)] and origin["rows"] == rows
        assert any(step["details"].get("method") == "derived_from_storage_and_loss_modulus"
                   for step in dataset["provenance"]["transforms"]) is derived
        assert dataset["columns"]["x"]["unit"] == "rad/s" and dataset["columns"]["y"]["unit"] == "Pa s"
    assert {path: path.read_bytes() for path in originals} == originals


def acquisition_book(tmp_path, name="acquisitions.xlsx", *, ambiguous=False):
    book = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TGA"
    for row in ([["Time", "Mass", "Weight"], ["s", "%", "%"], [0, 100, 70], [1, 99, 69]] if ambiguous
                else [["Time", "Mass"], ["s", "%"], [0, 100], [1, 99]]):
        sheet.append(row)
    sheet = workbook.create_sheet("Mechanical")
    for row in [["Time", "Storage Modulus"], ["s", "Pa"], [0, 10], [1, 20]]:
        sheet.append(row)
    workbook.save(book)
    return book


def table_request(sources):
    return {"version": 1, "action": "create", "sources": [str(source) for source in sources],
            "intent": {"x": {"quantity": "time"}, "y": {"quantity": "mass"}}}


def answer(result, **choices):
    return {"expected_question_id": result["question"]["question_id"], "confirm_plan": True, "choices": choices}


def test_unresolved_alternative_table_cannot_block_a_valid_explicit_coordinate(tmp_path):
    book = acquisition_book(tmp_path)
    value = table_request([book])["intent"]
    value["y"]["header"] = "Mass"
    plan = prepare_plan([book], value)
    assert [decision["field"] for decision in plan["decisions"]] == ["tables"]
    assert plan["decisions"][0]["options"] == {"source:0": ["table:0", "table:1"]}
    assert len(plan["sources"][0]["candidates"]) == 2


def test_table_answer_replans_only_selected_facts_and_keeps_stale_reply_safe(tmp_path, monkeypatch):
    from sciplot_core import task_control, task_semantic_creation as creation

    source = acquisition_book(tmp_path, ambiguous=True)
    original = source.read_bytes()
    calls = []
    monkeypatch.setattr(creation, "execute_managed", lambda current, wire: calls.append(wire) or {
        "status": "blocked", "ready_to_use": False, "next_step": {"action": "retry_same_request"}})
    task = tmp_path / "task"
    first = task_control.start_task(table_request([source]), task_dir=task)
    assert first["status"] == "needs_input" and not calls
    second = task_control.resume_task(task, answer(first, tables={"source:0": "table:0"}))
    assert second["status"] == "needs_input" and not calls
    assert second["question"]["question_id"] != first["question"]["question_id"]
    assert [decision["field"] for decision in second["question"]["plan"]["decisions"]] == ["y_header"]
    before = (task / "task.json").read_bytes()
    with pytest.raises(TaskControlError) as stale:
        task_control.resume_task(task, answer(first, tables={"source:0": "table:0"}))
    assert stale.value.reason_code == "stale_task_question"
    assert (task / "task.json").read_bytes() == before
    final = task_control.resume_task(task, answer(second, y_header="Mass"))
    assert len(calls) == 1
    assert calls[0]["data_binding"]["data_sources"][0]["table_selection"]["sheet"] == "TGA"
    assert final["source_plan"]["selection"]["selected_block_count"] == 1
    assert final["source_plan"]["selection"]["excluded"][0]["table_id"] == "table:1"
    task_control.resume_task(task, {"retry": True})
    assert len(calls) == 2 and calls[0] == calls[1]
    assert source.read_bytes() == original


def test_partial_multisource_table_answer_is_rejected_without_silently_selecting_rest(tmp_path):
    from sciplot_core import task_control

    first, second = acquisition_book(tmp_path), acquisition_book(tmp_path, "second.xlsx")
    task = tmp_path / "task"
    result = task_control.start_task(table_request([first, second]), task_dir=task)
    assert result["status"] == "needs_input"
    assert set(result["question"]["plan"]["decisions"][0]["options"]) == {"source:0", "source:1"}
    before = (task / "task.json").read_bytes()
    with pytest.raises(TaskControlError) as rejected:
        task_control.resume_task(task, answer(result, tables={"source:0": "table:0"}))
    assert rejected.value.reason_code == "invalid_task_response"
    assert (task / "task.json").read_bytes() == before


def test_complete_multisource_table_answer_binds_each_original_selected_table(tmp_path, monkeypatch):
    from sciplot_core import task_control, task_semantic_creation as creation

    first, second = acquisition_book(tmp_path), acquisition_book(tmp_path, "second.xlsx")
    calls = []
    monkeypatch.setattr(creation, "execute_managed", lambda current, wire: calls.append(wire) or {
        "status": "blocked", "ready_to_use": False, "next_step": {"action": "retry_same_request"}})
    task = tmp_path / "task"
    started = task_control.start_task(table_request([first, second]), task_dir=task)
    result = task_control.resume_task(task, answer(started, tables={"source:0": "table:0", "source:1": "table:0"}))
    assert len(calls) == 1
    assert [item["source_id"] for item in calls[0]["data_binding"]["data_sources"]] == ["source:0", "source:1"]
    assert all(item["table_selection"]["sheet"] == "TGA" for item in calls[0]["data_binding"]["data_sources"])
    assert result["source_plan"]["selection"]["selected_point_count"] == 4
    assert len(result["source_plan"]["selection"]["excluded"]) == 2


def test_legacy_saved_combined_question_can_answer_tables_and_coordinates_once(tmp_path, monkeypatch):
    from sciplot_core import task_control, task_semantic_creation as creation
    from openpyxl import load_workbook

    source = acquisition_book(tmp_path, ambiguous=True)
    workbook = load_workbook(source)
    workbook.remove(workbook["Mechanical"])
    alternate = workbook.create_sheet("Other mass acquisition")
    for row in [["Time", "Mass", "Weight"], ["s", "%", "%"], [0, 80, 60], [1, 79, 59]]:
        alternate.append(row)
    workbook.save(source)
    plan = prepare_plan([source], table_request([source])["intent"])
    plan["decisions"].append({"field": "y_header", "options": ["Mass", "Weight"]})
    monkeypatch.setattr(creation, "prepare_plan", lambda paths, value: plan)
    calls = []
    monkeypatch.setattr(creation, "execute_managed", lambda current, wire: calls.append(wire) or {
        "status": "blocked", "ready_to_use": False, "next_step": {"action": "retry_same_request"}})
    task = tmp_path / "legacy-task"
    started = task_control.start_task(table_request([source]), task_dir=task)
    result = task_control.resume_task(task, answer(started, tables={"source:0": "table:0"}, y_header="Mass"))
    assert len(calls) == 1 and result["status"] == "blocked"
    assert "question" not in result
    assert calls[0]["data_binding"]["provenance"]["semantic_task"]["choices"] == {
        "tables": {"source:0": "table:0"}, "y_header": "Mass"}
